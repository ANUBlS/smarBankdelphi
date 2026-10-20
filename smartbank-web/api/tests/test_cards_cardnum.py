"""Cards ▸ Card Number Change (cardnumbersource.pas) and Commission Our POS-ONUS (PostCommisionOnusImpl.pas)."""
import base64
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.db import mock as mock_db
from app.db import mock_cards as MC

client = TestClient(app)


def P(path, body=None, code=200):
    r = client.post("/api/cards" + path, json=body or {})
    assert r.status_code == code, r.text
    return r.json() if r.headers.get("content-type", "").startswith("application/json") else r


def login():
    mock_db.reset(); MC.reset()
    assert client.post("/api/login", json={"username": "ADMIN", "password": "Admin#2026"}).status_code == 200


def line(ref, card, name, card2, acc):
    s = list(" " * 140)
    for pos, v in ((1, ref), (34, card), (51, name), (89, card2), (106, acc)):
        s[pos - 1:pos - 1 + len(v)] = v
    return "".join(s).rstrip()


def test_cardnum_auto():
    login()
    v = P("/cardnum/open")
    assert v["page"] == "auto" and v["grid1"] == [] and [r["referance"] for r in v["grid2"]] == ["700400"]
    data = "\r\n".join([line("700501", "5239820000005555", "ISMAYILOV TURAL", "5239820000005599", "38110019440000000005"),
                        line("799999", "4169730000001111", "NOBODY", "4169730000001199", "38110019440000000001")]).encode("cp1254")
    P("/cardnum/upload", {"name": "CHG.txt", "b64": base64.b64encode(data).decode()})
    r = P("/cardnum/tobase")
    assert r["lines"] == 2 and [(x["referance"], x["cardnumber"], x["name"], x["status"]) for x in r["grid1"]] == [
        ("700501", "5239820000005599", "ISMAYILOV TURAL", 1), ("799999", "4169730000001199", "NOBODY", 3)]
    ids = [x["id"] for x in r["grid1"]]
    n = P("/cardnum/change", {"ids": ids, "answers": ["no", "no"]})
    assert n["done"] == [] and n["messages"] == []
    c = P("/cardnum/change", {"ids": ids, "answers": ["all"]})                # Yes to all
    assert c["done"] == ids and c["messages"] == ["Kart nömrələri dəyişildi!"] and len(c["grid2"]) == 2
    assert MC._db().execute("select cardnumber from card_clients where id = 5").fetchone()[0] == "5239820000005599"
    d = P("/cardnum/delete", {"current": ids[1], "answers": ["yes"]})        # nothing selected -> the current row
    assert d["grid1"] == [] and d["messages"] == ["Sətir silindi!"]
    x = P("/cardnum/excel/grid2")
    assert load_workbook(io.BytesIO(x.content)).sheetnames == ["KART NOMRELERI"]


def test_cardnum_manual():
    login(); P("/cardnum/open")
    m = P("/cardnum/tab", {"page": "manual"})
    assert m["boxes"]["DateTimePicker7"] == "TODAY" and m["boxes"]["Edit4"] == ""
    e = P("/cardnum/edit1", {"edit1": "700501"})
    assert e["boxes"]["Edit15"] == "Tural" and e["boxes"]["DateTimePicker7"] == "1985-01-01"   # client2 row found
    assert e["boxes"]["Edit26"] == "700501" and e["edit2"] == ""
    assert P("/cardnum/edit1", {"edit1": "NOPE"})["boxes"]["Edit4"] == ""
    s0 = P("/cardnum/save", {"edit1": "700501", "edit2": "", "yes": True})
    assert "messages" not in s0                                                # Edit2 empty -> nothing
    s = P("/cardnum/save", {"edit1": "700501", "edit2": "5239820000005588", "yes": True})
    assert s["messages"] == ["Kart nömrəsi bazaya qeyd olundu!"] and s["boxes"]["Edit24"] == "5239820000005588" and s["edit2"] == "5239820000005588"
    assert P("/cardnum/tab", {"page": "auto"})["page"] == "auto"


def test_posonus():
    login()
    v = P("/posonus/open")
    assert v["date"] == "2026-10-12" and [r["id"] for r in v["grid1"]] == [1, 2] and [r["id"] for r in v["grid2"]] == [3]
    assert P("/posonus/refresh", {"date": "2026-10-11"})["grid1"][0]["id"] == 4
    P("/posonus/refresh", {"date": "2026-10-12"})
    c = P("/posonus/confirm", {"ids": [1, 2], "answers": ["cancel", "yes"]})   # Cancel skips row 1; Yes processes row 2 (no more questions)
    assert c["done"] == [2] and c["messages"] == ["1 sayda əməliyyat təsdiqləndi!"] and [r["id"] for r in c["grid2"]] == [2, 3]
    r = P("/posonus/read")
    assert len(r["grid1"]) == 3
    d = P("/posonus/delete", {"grid": "grid2", "ids": [2, 3], "answers": ["yes"]})
    assert d["messages"] == ["2 sayda seçilmiş əməliyyat silindi!"] and d["grid2"] == []
    d1 = P("/posonus/delete", {"grid": "grid1", "current": 1, "answers": ["yes"]})
    assert d1["messages"] == ["1 sayda seçilmiş əməliyyat silindi!"]
    P("/posonus/refresh", {"date": "2026-10-01"})
    assert P("/posonus/confirm", {"ids": [], "answers": []})["messages"] == ["Hec bir melumat yoxdur."]
