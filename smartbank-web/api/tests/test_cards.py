"""Packages/Cards_New/cards_azericard_turan_version (CARDS, TMAINFORM) step 1 - the card-holder list (customers.pas Tclients).
Expected values traced by hand from customers.pas / Unit18.dfm and fixtures/cards.json."""
import datetime

from fastapi.testclient import TestClient

from app.main import app
from app.db import mock as mock_db
from app import cards as C
from app import cards_sql as Q
from app import tdataset_filter as TF

client = TestClient(app)
B = "/api/cards"


def login(u="ADMIN", p="Admin#2026"):
    mock_db.reset()
    assert client.post("/api/login", json={"username": u, "password": p}).status_code == 200


def P(path, body=None, code=200):
    r = client.post(B + path, json=body or {})
    assert r.status_code == code, r.text
    return r.json()


def test_sql_verbatim_lines():
    assert Q.LINES_CLIENTS[0] == "SELECT   s.ROWID," and Q.LINES_CLIENTS[19] == "where s.accid=b.id"
    assert len(Q.LINES_CLIENTS) == 21                                           # cards_azericard: no SQL[25] line
    assert Q.LINES_CLIEXCEL[19] == "and s.multiacc is null order by s.rbsnumber"
    assert Q.LINES_SETTINGS == ["select s.rowid, s.* from card_settings s"]
    assert Q.SOURCES["ClientListSource"] == "Clients"


def test_open_admin_all_filials():
    login()
    r = P("/clients/open")
    assert r["count"] == 5 and r["dateclose"] == 30
    a = {x["id"]: x for x in r["rows"]}
    assert a[1]["fil"] == "Baş ofis" and a[3]["fil"] == "Nərimanov filialı" and a[1]["creat"] == "ADMIN" and a[2]["modi"] == "TELLER"
    assert a[1]["cardnumber"] == "4169730000001111" and a[1]["old_acc_no"] == "38110000000001"
    assert r["lookups"]["cardstatus"] == ["ACTIVE", "BLOCKED", "NEW"] and r["lookups"]["company"] == ["TURANBANK", "AZERSUN"]
    assert r["hidden"] == [] and "VERI" not in r["cols"] and r["cols"][2] == "MULTIACC"


def test_open_teller_sees_all_filials():
    """cards_azericard FormActivate has no CARD_FIL_VIEW branch (Turan-only): every user gets all rows."""
    login("TELLER", "Teller#2026")
    r = P("/clients/open")
    assert sorted(x["id"] for x in r["rows"]) == [1, 2, 3, 4, 5]


def test_excel_export_cliexcel():
    import io
    from openpyxl import load_workbook
    login(); P("/clients/open")
    P("/clients/edit-filter", {"field": "FNAME", "text": "R"})               # Panel4 filter does not reach cliexcel
    r = client.post(B + "/clients/excel")
    assert r.status_code == 200
    sh = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in sh[1]][:4] == [" #", "RBSNUMBER", "CARDHOLDER", "CARDNUMBER"] and sh.max_column == 28
    assert sh["A1"].font.bold
    rbs = [sh.cell(i, 2).value for i in range(2, sh.max_row + 1)]
    # multiacc row (id 2) left out; Oracle puts the NULL RBSNUMBER last, the SQLite mock first - compare without it
    assert [x for x in rbs if x] == ["38110019440000000001", "38110019440000000004", "38110019840000000003"] and len(rbs) == 4
    P("/clients/filter", {"checks": [False, False, True, False, False, False], "d1": "2023-01-01", "d2": "2027-12-31", "valuta": "USD"})
    sh = load_workbook(io.BytesIO(client.post(B + "/clients/excel").content)).active
    assert all(sh.cell(i, 7).value == "USD" for i in range(2, sh.max_row + 1))


def test_open_needs_card_clients():
    login("ACCOUNTANT", "Accountant#2026")
    r = client.post(B + "/clients/open")
    assert r.status_code in (403, 400)


def test_edit_filter_prefix_case_sensitive():
    login(); P("/clients/open")
    r = P("/clients/edit-filter", {"field": "FNAME", "text": "R"})
    assert r["filter"] == "FNAME = 'R*'" and [x["id"] for x in r["rows"]] == [3]
    assert P("/clients/edit-filter", {"field": "FNAME", "text": "r"})["count"] == 0          # FilterOptions = []
    r = P("/clients/edit-filter", {"field": "CARDNUMBER", "text": "5239"})                  # replaces the FNAME filter
    assert sorted(x["id"] for x in r["rows"]) == [4, 5]
    r = P("/clients/edit-filter", {"field": "CREDIT_LIM", "text": "5"})                     # numeric field vs '5*'
    assert r["filter_error"] and r["count"] == 0
    P("/clients/edit-filter", {"field": "PASSPORT", "text": "x"}, 400)


def test_refresh_hides_null_rbs():
    login(); P("/clients/open")
    r = P("/clients/refresh")
    assert r["filter"] == "RBSNUMBER = '*'" and sorted(x["id"] for x in r["rows"]) == [1, 2, 3, 4]


def test_panel_filter_combinations():
    assert C.panel5_filter([False, False, True, False, True, False], "", "", "", "AZN", "", "ACTIVE", "") == \
        "(CURRENCY = 'AZN') and (CARDSTATUS = 'ACTIVE')"
    # 1+3+5: Delphi's table has no f3 for this set - the status condition is dropped (kept)
    assert C.panel5_filter([True, False, True, False, True, False], "01.01.2024", "31.12.2027", "", "AZN", "", "ACTIVE", "") == \
        "(DATEOPEN >= '01.01.2024') and (EXPIRY <= '31.12.2027') and (CURRENCY = 'AZN')"
    login(); P("/clients/open")
    r = P("/clients/filter", {"checks": [True, False, False, False, False, True], "d1": "2023-01-01", "d2": "2027-12-31",
                              "company": "TURANBANK"})
    assert r["filter"] == "(DATEOPEN >= '01.01.2023') and (EXPIRY <= '31.12.2027') and (COMPANY = 'TURANBANK')"
    assert [x["id"] for x in r["rows"]] == [1]


def test_sort_toggle_and_delete():
    login(); P("/clients/open")
    assert [x["id"] for x in P("/clients/sort", {"field": "LNAME"})["rows"]] == [2, 3, 1, 4, 5]
    assert [x["id"] for x in P("/clients/sort", {"field": "LNAME"})["rows"]] == [5, 4, 1, 3, 2]
    assert P("/clients/sort", {"field": "ROWID"})["count"] == 5                              # not a grid column: no change
    r = P("/clients/delete", {"id": 2, "confirmed": False})
    assert r["count"] == 5
    r = P("/clients/delete", {"id": 2, "confirmed": True})
    assert sorted(x["id"] for x in r["rows"]) == [1, 3, 4, 5]


def test_filter_parser():
    rows = [{"A": "abc", "D": datetime.datetime(2025, 1, 1), "N": None}]
    assert TF.apply(rows, "A = 'ab*'", {"A", "D", "N"}) == rows
    assert TF.apply(rows, "(D >= '01.01.2025') and (A <> 'x')", {"A", "D", "N"}) == rows
    assert TF.apply(rows, "N = '*'", {"A", "D", "N"}) == []
    try:
        TF.apply(rows, "A = 'x", {"A"}); assert False
    except TF.FilterError:
        pass
