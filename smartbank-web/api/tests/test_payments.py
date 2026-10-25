"""Payments 'Köçürmələr' (payments.py, pay_sql.py, db/mock_pay.py)."""
import pytest
from fastapi.testclient import TestClient
from app import main, sessions, pay_sql, payments
from app.db import mock, mock_pay


@pytest.fixture(autouse=True)
def fresh():
    mock.reset(); sessions.attempts._f.clear(); sessions.store._s.clear(); yield


def cl(u="ADMIN", p="Admin#2026"):
    c = TestClient(main.app); assert c.post("/api/login", json={"username": u, "password": p}).status_code == 200; return c


def test_sql_verbatim_and_substitution():
    L = pay_sql.SQL_PAY_ODSPAYMENTS.split("\n")
    assert L[33] == "AND    (trunc(a.susdat)) >=:DAT1 AND  pay_date <=:DAT2 and a.status not in (0,9,10)"
    assert L[39] == "" and L[44] == " order by :sort"
    s = payments.payments_sql("STATUS", 2).split("\n")
    assert s[39] == "where filial = '2'" and s[44] == " order by STATUS" and s[:39] == L[:39]
    assert pay_sql.SQL_PAY_DSFILIALS == "select * from Filial order by id"
    assert pay_sql.SQL_PAY_ODS_REP.endswith("AND    a.id=:v_id")


def test_init_open_sort_filial():
    c = cl()
    i = c.get("/api/payments/init").json()
    assert i["tek_dat"] == "2026-10-12" and [f["id"] for f in i["filials"]] == [1, 2, 3]
    ids = [r["id"] for r in i["rows"]]
    assert ids == [1, 2, 3, 4, 5, 6, 7, 8, 9]          # status 9 / type 44 / trans_id null / other day excluded; order by 1
    r6 = next(r for r in i["rows"] if r["id"] == 6)
    assert r6["alan_acc"] == "AZ99SMRT33010018400099999"  # type 7, to_accid -1 -> decode(.. null) -> receiver_acc
    r5 = next(r for r in i["rows"] if r["id"] == 5)
    assert r5["alan_name"] == "Həsənov Elçin (USD)" and r5["tr_type"] == "Konvertasiya"
    o = c.post("/api/payments/open", json={"dat1": "2026-10-09", "dat2": "2026-10-12", "sort": "AMOUNT"}).json()["rows"]
    assert [r["amount"] for r in o] == sorted(r["amount"] for r in o) and 13 in [r["id"] for r in o]
    f = c.post("/api/payments/open", json={"dat1": "2026-10-12", "dat2": "2026-10-12", "filial": 2}).json()["rows"]
    assert {r["filial"] for r in f} == {2}
    assert c.post("/api/payments/open", json={"dat1": "2026-10-12", "dat2": "2026-10-12", "sort": "1;drop"}).status_code == 400
    assert c.post("/api/payments/open", json={"dat1": "x", "dat2": "2026-10-12"}).status_code == 422


def test_check_status_calls_and_doc_path():
    c = cl()
    r = c.post("/api/payments/check_status", json={"trans_ids": ["IB000001", "IB000002"]}).json()
    assert r["calls"] == 3 and mock_pay.CALLS == [("IB000001", 3), ("IB000002", 3), ("IB000002", 3)]   # j always 3, last twice
    p = c.post("/api/payments/doc_path", json={"id": 2, "ext": "pdf"}).json()["path"]
    assert p == "\\\\10.0.11.90\\adocs\\7101-2.pdf"
    assert c.post("/api/payments/doc_path", json={"id": 2, "ext": "exe"}).status_code == 422


def test_excel_filtered():
    c = cl()
    x = c.post("/api/payments/excel", json={"dat1": "2026-10-12", "dat2": "2026-10-12",
                                           "filter": {"types": [1, 2, 0, 0, 0, 0, 0, 0], "statuses": [1, 2, 3]}})
    assert x.status_code == 200
    from openpyxl import load_workbook
    import io
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.title == "Transactions" and ws.max_row == 3       # header + ids 1, 2 (id 9 has status 4)


def test_rights_denied():
    c = cl("TELLER", "Teller#2026")
    assert c.get("/api/payments/init").status_code == 403
    assert c.post("/api/payments/check_status", json={"trans_ids": []}).status_code == 403
