"""Cards ▸ Excel yükləmələr sub forms (cards_xls): Unitsalary, EXCELNEWCARD, PROLONGATIONUNIT, accountcloseunit."""
import base64
import datetime as dt
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app
from app.db import mock as mock_db
from app.db import mock_cards as MC
from app import cards_xls as X

client = TestClient(app)


def P(path, body=None, code=200):
    r = client.post("/api/cards" + path, json=body or {})
    assert r.status_code == code, r.text
    return r.json() if r.headers.get("content-type", "").startswith("application/json") else r


def login():
    mock_db.reset(); MC.reset()
    assert client.post("/api/login", json={"username": "ADMIN", "password": "Admin#2026"}).status_code == 200
    P("/excelforms/open")


def wb(cells, sheet="CARDS"):
    w = Workbook(); s = w.active; s.title = sheet
    for (r, c), v in cells.items():
        s.cell(r, c, v)
    b = io.BytesIO(); w.save(b)
    return base64.b64encode(b.getvalue()).decode()


def test_requires_parent_form():
    mock_db.reset(); MC.reset()
    client.post("/api/login", json={"username": "ADMIN", "password": "Admin#2026"})
    P("/xls/salary/open", code=409)


def test_var_to_str_and_fields():
    assert X.var_to_str(150.5) == "150,5" and X.var_to_str(12.0) == "12" and X.var_to_str(None) == ""
    assert X.var_to_str(4169738812345678.0) == "4,16973881234568E15"           # FloatToStr: 15 significant digits
    assert X.var_to_str(dt.datetime(2026, 11, 30)) == "30.11.2026" and X.var_to_str(True) == "True"
    assert X.set_as_string("AMMOUNT", "150,5") == 150.5 and X.set_as_string("AMMOUNT", "") is None
    assert X.set_as_string("DAT", "30.11.2026") == dt.datetime(2026, 11, 30)
    for f, v, m in (("AMMOUNT", "abc", "not a valid floating point value"), ("DAT", "2026-11-30", "not a valid date and time")):
        try:
            X.set_as_string(f, v); raise AssertionError
        except ValueError as e:
            assert m in str(e)


def test_salary():
    login()
    v = P("/xls/salary/open")                                           # FormActivate: both open
    assert v["grid1"] == [] and [r["cardholder"] for r in v["grid2"]] == ["ALIYEVA NIGAR"]
    cells = {(6, 3): "OKT-2026", (10, 2): "MAMMADOV ELCHIN", (10, 3): "38110019440000000001", (10, 4): 120.5, (10, 5): "AZN", (10, 6): "38010019440000000001",
             (11, 2): " QULIYEV ", (11, 4): 80, (11, 7): "ignored"}
    r = P("/xls/salary/tobase", {"name": "s.xlsx", "b64": wb(cells)})
    assert r["messages"] == ["2 record is imported!"]
    g = {x["cardholder"]: x for x in r["grid1"]}
    assert g["MAMMADOV ELCHIN"]["ammount"] == 120.5 and g["MAMMADOV ELCHIN"]["status"] == 1
    assert g[" QULIYEV "]["status"] == 3                                 # no Trim in Unitsalary; saltr (mock) marks the empty RBSNUMBER
    s = P("/xls/salary/sort", {"grid": "grid1", "field": "AMMOUNT"})
    assert [x["ammount"] for x in s["grid1"]] == [80, 120.5]
    n = P("/xls/salary/confirm", {"yes": False})
    assert len(n["grid1"]) == 2
    y = P("/xls/salary/confirm", {"yes": True})
    assert [x["status"] for x in y["grid1"]] == [3] and len(y["grid2"]) == 2
    d = P("/xls/salary/delete")
    assert d["grid1"] == []
    x = P("/xls/salary/excel", {"grid": "grid2", "field": ""})
    assert load_workbook(io.BytesIO(x.content)).sheetnames == ["salary"]
    assert P("/xls/salary/close")["ok"]


def test_salary_no_file_and_errors():
    login(); P("/xls/salary/open")
    r = P("/xls/salary/tobase", {})                                       # dialog cancelled: saltr('') + Refresh only
    assert r["messages"] == [] and r["grid1"] == []
    e = P("/xls/salary/tobase", {"name": "s.xlsx", "b64": wb({(10, 2): "A", (10, 4): "x1", (11, 2): "B"})})
    assert "'x1' is not a valid floating point value" in e["error"] and e["messages"] == []
    P("/xls/salary/tobase", {"name": "s.xlsx", "b64": wb({(10, 2): "A"}, sheet="Sheet1")}, code=400)
    P("/xls/salary/tobase", {"name": "s.txt", "b64": base64.b64encode(b"hello").decode()}, code=400)


def test_newcard():
    login()
    v = P("/xls/newcard/open")
    assert v["date"] == dt.date.today().isoformat() and v["grid1"] == [] and len(v["grid2"]) == 1
    cells = {(1, 2): "NAME"}
    for i, val in enumerate([" LEYLA ", "KARIMOVA", "01.02.1995", "AZE3000001", "RESIDENT", "MOTHER", "0505550002", "VISA CLASSIC",
                             "38110019440000000006", "Baku", "TURANBANK", 501]):
        cells[(2, 2 + i)] = val
    r = P("/xls/newcard/tobase", {"name": "n.xlsx", "b64": wb(cells)})
    assert r["messages"] == ["Tarixi düz seçin!"] and r["grid1"] == []    # DateTimePicker1 = today -> not > DATE
    tm = (dt.date.today() + dt.timedelta(days=1)).isoformat()
    r = P("/xls/newcard/tobase", {"name": "n.xlsx", "b64": wb(cells), "date": tm})
    assert r["messages"] == ["1 record is imported!"]
    row = r["grid1"][0]
    assert row["name"] == "LEYLA" and row["id"] == 501 and row["company"] == "TURANBANK" and row["status"] == 1   # Trim; ID from column 13
    c = P("/xls/newcard/confirm")
    assert c["grid1"] == [] and len(c["grid2"]) == 2


def test_prolong_and_accclose():
    login()
    v = P("/xls/prolong/open")                                           # no FormActivate: closed until something opens PROLONG
    assert v["grid1"] is None and v["grid2"] is None
    cells = {(1, 1): "CARDNUMBER", (1, 2): "DAT", (2, 1): "4169730000001111", (2, 2): dt.datetime(2029, 3, 31), (3, 2): "31.12.2028"}
    r = P("/xls/prolong/tobase", {"name": "p.xlsx", "b64": wb(cells)})
    assert r["messages"] == ["2 record is imported!"] and r["grid2"] is None     # PROLONG.Refresh twice, PROLONG1 not refreshed
    assert [x["dat"] for x in r["grid1"]] == ["2029-03-31", "2028-12-31"] and [x["status"] for x in r["grid1"]] == [1, 3]
    c = P("/xls/prolong/confirm")
    assert [x["status"] for x in c["grid1"]] == [3] and len(c["grid2"]) == 2
    v2 = P("/xls/prolong/open")                                          # PROLONG / PROLONG1 stay open in the DataMod
    assert v2["grid1"] is not None and v2["grid2"] is not None
    a = P("/xls/accclose/open")
    assert a["grid1"] is None
    r = P("/xls/accclose/tobase", {"name": "a.xlsx", "b64": wb({(2, 1): "4169730000003333", (2, 2): "", (3, 1): "5239820000005555"})})
    assert r["messages"] == ["2 record is imported!"] and len(r["grid2"]) == 1
    e = P("/xls/accclose/tobase", {"name": "a.xlsx", "b64": wb({(2, 1): "4169730000003333", (2, 2): "OK"})})
    assert "'OK' is not a valid floating point value" in e["error"]          # Fields[1] = STATUS
    d = P("/xls/accclose/delete")
    assert d["grid1"] == []
    assert P("/xls/accclose/excel", {"grid": "grid1", "field": ""}).content[:2] == b"PK"
