"""CashDeskLast (TFrmCashDesk) - api/app/cashdesk.py on the mock (api/app/db/mock_cd.py)."""
import base64, io
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import main, sessions, cashdesk as CD
from app.db import mock, mock_cd


@pytest.fixture(autouse=True)
def fresh():
    mock.reset(); sessions.attempts._f.clear(); sessions.store._s.clear(); yield


def cl(u="ADMIN", p="Admin#2026"):
    c = TestClient(main.app); assert c.post("/api/login", json={"username": u, "password": p}).status_code == 200; return c


def post(c, url, code=200, **b):
    r = c.post(url, json=b); assert r.status_code == code, r.text; return r.json()


U = "/api/cashdesk/"
D = dict(beg="2026-10-12", end="2026-10-12")


def test_init_lists_and_date():
    c = cl(); o = c.get(U + "init").json()
    assert o["date"] == "2026-10-12"                                   # GetPresentDate = SYSDAT.TEK_DAT
    assert o["currencies"][0]["name"] == "Bütün valyutalar" and len(o["currencies"]) == 4
    assert o["filials"][0]["name"] == "Bütün filiallar" and [f["id"] for f in o["filials"][1:]] == [101, 102]
    assert o["cash_symbols"][0] == "Bütün" and o["operators"][0]["name"] == "Butun operatorlar"
    assert o["grid_cols"] == CD.GRID_COLS
    assert set(o["flags"].values()) <= {"0", "1"} and len(o["flags"]) == 11


def test_binds_follow_tbrefreshclick():
    L = {"rate": [0, 840, 978], "filials": [{"ID": 101}, {"ID": 102}]}
    b = CD.binds(CD.Filt(**D), L)
    assert b["p_currencyid"] is None and b["p_cashsymbol"] is None and b["p_userinsert"] is None and b["p_usercommit"] is None
    assert b["p_debitbranch"] is None and b["p_creditbranch"] is None and b["p_orderby"] is None
    assert b["p_transtype"] == 1 and b["p_balance"] == 1
    b = CD.binds(CD.Filt(**D, currency=2, cash_symbol="10", operator="TELLER", committer="ADMIN", debit_branch=1,
                         credit_branch=2, transtype=2, balance=3, orderby="DAT"), L)
    assert b["p_currencyid"] == 840 and b["p_cashsymbol"] == "10" and b["p_userinsert"] == "TELLER"
    assert b["p_usercommit"] == "ADMIN" and b["p_debitbranch"] == "101" and b["p_creditbranch"] == "102"
    assert b["p_transtype"] == 2 and b["p_balance"] == 4 and b["p_orderby"] == "DAT"


def test_refresh_filters_and_orderby_whitelist():
    c = cl()
    rows = post(c, U + "refresh", **D)["rows"]
    assert len(rows) == 8 and all(r["dat"].startswith("2026-10-12") for r in rows)
    all_cur = len(rows)
    one = post(c, U + "refresh", **D, currency=2)["rows"]
    assert 0 < len(one) < all_cur
    lazy = post(c, U + "refresh", beg="2026-10-01", end="2026-10-31", transtype=2)["rows"]
    assert {r["id"] for r in lazy} == {9201, 9202}
    post(c, U + "refresh", **D, orderby="LAMMOUNT")
    assert post(c, U + "refresh", 400, **D, orderby="1; drop table x")["detail"]["key"] == "error.generic"
    assert post(c, U + "refresh", 400, **D, currency=99)["detail"]["key"] == "error.generic"
    assert post(c, U + "refresh", 422, **D, transtype=3)


def test_choose_cash_symbol():
    c = cl()
    assert post(c, U + "symbol", id=9101, symbol="")["changed"] is False
    assert post(c, U + "symbol", id=9101, symbol="42")["changed"] is True
    r = next(x for x in post(c, U + "refresh", **D)["rows"] if x["id"] == 9101)
    assert str(r["kassa_b"]) == "42"
    assert post(c, U + "symbol", 400, id=9101, symbol="abc")["detail"]["key"] == "cashdesk.wrong_args"


def test_packno_table_and_quote():
    c = cl()
    post(c, U + "packno", id=9101, packno="77", transtype=1)
    assert mock_cd._db().execute("select NO_PROV from provarch where ID=9101").fetchone()[0] == "77"
    post(c, U + "packno", id=9201, packno="5", transtype=2)
    assert mock_cd._db().execute("select NO_PROV from lazytransaction where ID=9201").fetchone()[0] == "5"
    assert post(c, U + "packno", 400, id=9101, packno="1' or '1'='1")["detail"]["key"] == "cashdesk.err_packno_quote"


def test_trans_commit_d1_rule():
    c = cl()
    assert post(c, U + "commit", 400, id=9202, d1="1")["detail"]["key"] == "cashdesk.not_checked"
    post(c, U + "commit", id=9201)
    db = mock_cd._db()
    assert db.execute("select count(*) from lazytransaction where ID=9201").fetchone()[0] == 0
    assert db.execute("select count(*) from provarch where ID=9201").fetchone()[0] == 1
    r = c.post(U + "commit", json={"id": 123456}); assert r.status_code >= 400


def test_cashinfo_enable_rules():
    c = cl()
    o = c.get(U + "cashinfo/9104").json()
    assert o["info"]["issuedby"] == "Həsənova Leyla Vaqif qızı" and o["browse"]
    assert post(c, U + "cashinfo/9104", 400, action="insert", issuedby="x")["detail"]["key"] == "error.generic"
    post(c, U + "cashinfo/9104", action="update", issuedby="Yeni", person_doc="Pasport", seria="", docnumber="1", fromwhen="")
    o = c.get(U + "cashinfo/9104").json()["info"]
    assert o["issuedby"] == "Yeni" and o["seria"] is None                  # '' -> NULL
    post(c, U + "cashinfo/9104", action="delete")
    assert c.get(U + "cashinfo/9104").json()["info"] is None
    assert post(c, U + "cashinfo/9104", 400, action="update")["detail"]["key"] == "error.generic"
    post(c, U + "cashinfo/9104", action="insert", issuedby="Z")
    assert c.get(U + "cashinfo/9104").json()["info"]["issuedby"] == "Z"
    assert post(c, U + "cashinfo/9104", 400, action="drop")


def test_settings_19_args():
    c = cl()
    body = dict(natbalance="1001", forbalance="1002", description_az="Bank AZ", description_eng="Bank EN", signatyre="S",
                responsible_person="R", handover="1", chief="0", p_operator="1", cashier="1", amount="1", chief_="1",
                accountant="0", operator_="1", cashier_="1", controller="1", controller_="0", operator_title="Op",
                accountant_title="Acc")
    s = post(c, U + "settings", **body)["settings"]
    assert s["description_az"] == "Bank AZ" and s["operatortitle"] == "Op" and s["accountanttitle"] == "Acc"
    assert str(s["chief"]) == "0" and str(s["controller_out"]) == "0"
    f = c.get(U + "init").json()["flags"]
    assert f["chief"] == "0" and f["accountant"] == "0" and f["controller_"] == "0" and f["handover"] == "1"
    assert len(CD.SETTINGS_ARGS) == 19


def test_photo_rules():
    c = cl()
    bmp = b"BM" + b"\0" * 60
    s = post(c, U + "settings/photo", name="logo.bmp", data=base64.b64encode(bmp).decode())["settings"]
    assert s["has_photo"] is True
    r = c.get(U + "settings/photo"); assert r.status_code == 200 and r.content == bmp
    assert post(c, U + "settings/photo", 400, name="x.jpg", data="")["detail"]["key"] == "cashdesk.err_picture_ext"
    assert post(c, U + "settings/photo", 400, name="x.bmp", data=base64.b64encode(b"nope").decode())["detail"]["key"] == "cashdesk.err_bitmap"
    assert c.get(U + "settings/photo").status_code == 200                 # failed loads change nothing
    assert post(c, U + "settings/photo", name="", data="")["settings"]["has_photo"] is False   # cancelled dialog clears
    post(c, U + "settings/photo", name="logo.bmp", data=base64.b64encode(bmp).decode())
    assert post(c, U + "settings/photo", name="a.ico", data="AAAA")["settings"]["has_photo"] is False
    post(c, U + "settings/photo", name="logo.bmp", data=base64.b64encode(bmp).decode())
    assert post(c, U + "settings/photo/clear")["settings"]["has_photo"] is False
    assert c.get(U + "settings/photo").status_code == 404


def test_excel():
    c = cl()
    r = c.post(U + "excel", json=D); assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content)); ws = wb["Cashdesk"]
    assert ws.max_row == 9 and ws.cell(1, 1).font.bold


def test_rights():
    t = cl("TELLER", "Teller#2026")
    assert t.get(U + "init").status_code == 200 and t.get("/api/cashsymbols").status_code == 200
    a = cl("ACCOUNTANT", "Accountant#2026")
    assert a.get(U + "init").status_code == 403 and a.get("/api/cashsymbols").status_code == 403
