"""Cards ▸ Application (cards_azericard applicationpas.pas TApplications) step 3a. Expected values traced from the .pas/.dfm and
fixtures/cards.json (card_client_app rows 9201-9206, TRANDATE = the server date)."""
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.db import mock as mock_db
from app.db import mock_cards as MC
from app import cards_app as A

client = TestClient(app)
B = "/api/cards/applications"


def login(u="ADMIN", p="Admin#2026"):
    mock_db.reset(); MC.reset()
    assert client.post("/api/login", json={"username": u, "password": p}).status_code == 200


def P(path, body=None, code=200):
    r = client.post(B + path, json=body or {})
    assert r.status_code == code, r.text
    return r.json()


def test_open_split_and_status_order():
    login()
    r = P("/open")
    # urgent: URG = 1, status in (1,2,3), rbsnumber = multi_acc or multi_acc = 'n/a'; FormShow Button6 -> ORDER BY STATUS
    assert [x["id"] for x in r["urgent"]] == [9202, 9201]
    # simple: URG = 0; 9205 status 4 and 9206 rbsnumber <> multi_acc are out
    assert [x["id"] for x in r["simple"]] == [9204, 9203]
    assert r["urgent"][0]["emeliyyat"] == "CREDIT LIMIT" and r["urgent"][0]["fil"] == "Baş ofis" and r["urgent"][0]["creater"] == "TELLER"
    assert r["cols"][1] == "" and r["status_img"] == {"1": 25, "2": 3, "3": 26, "4": 13}


def test_refresh_other_day_and_no_order():
    login(); P("/open")
    import datetime
    r = P("/refresh", {"date": (datetime.date.today() - datetime.timedelta(days=1)).isoformat()})
    assert r["urgent"] == [] and r["simple"] == []
    r = P("/refresh", {"date": datetime.date.today().isoformat()})
    assert sorted(x["id"] for x in r["urgent"]) == [9201, 9202]


def test_title_sort_and_blank_lines():
    login(); P("/open")
    r = P("/sort", {"grid": "simple", "field": "ID"})
    assert [x["id"] for x in r["simple"]] == [9203, 9204]
    st = None
    P("/sort", {"grid": "simple", "field": ""})                        # no FieldName (status icon column): Exit
    P("/sort", {"grid": "nope", "field": "ID"}, 400)


def test_excel_and_rights():
    login(); P("/open")
    x = client.post(B + "/excel", json={"grid": "urgent"})
    assert x.status_code == 200
    sh = load_workbook(io.BytesIO(x.content)).active
    assert sh.title == "Application" and sh.max_row == 3 and "EMELIYYAT" in [c.value for c in sh[1]]


def test_not_opened():
    login()
    P("/refresh", {"date": "2026-01-01"}, 409)


def test_sql_lines():
    assert A.DS["urgent"][17].endswith("a.URG=1 and  ( a.rbsnumber=a.multi_acc or a.multi_acc='n/a' )")
