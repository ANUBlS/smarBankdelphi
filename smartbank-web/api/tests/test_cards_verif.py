"""Cards ▸ Verifikasiya (cards_azericard unitverification.pas). urgent2 / simple2 = STATUS 4 rows of :dat1 (fixture 9205: URG 0)."""
import datetime as dt
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.db import mock as mock_db
from app.db import mock_cards as MC

client = TestClient(app)
B = "/api/cards/verification"


def login():
    mock_db.reset(); MC.reset()
    assert client.post("/api/login", json={"username": "ADMIN", "password": "Admin#2026"}).status_code == 200


def P(path, body=None, code=200):
    r = client.post(B + path, json=body or {})
    assert r.status_code == code, r.text
    return r.json()


def add(i, urg=1):
    today = dt.datetime.combine(dt.date.today(), dt.time()).isoformat()
    MC._db().execute("insert into card_client_app (ID, MESSTYPE, STATUS, URG, TRANDATE, FILIAL, OP_ID, RBSNUMBER, MULTI_ACC) "
                     "values (?, 1, 4, ?, ?, 101, 2, 'A', 'A')", [i, urg, today])


def ids(r, g):
    return [x["id"] for x in r[g]]


def test_open_refresh():
    login()
    r = P("/open")
    assert ids(r, "simple") == [9205] and r["urgent"] == [] and len(r["cols"]) == 27 and r["cols"][1] == ""
    assert r["simple"][0]["emeliyyat"] == "NEW" and r["msg_row"] == "Əmaliyyat verifikasiya edilsin? №"
    r = P("/refresh", {"date": "2020-01-01"})
    assert r["simple"] == [] and r["urgent"] == []


def test_verify_loop_answers():
    login()
    for i in (9401, 9402, 9403, 9404):
        add(i)
    r = P("/open")
    assert ids(r, "urgent") == [9401, 9402, 9403, 9404]
    # 9401 No, 9402 Yes, 9403 YesToAll+Cancel (Continue, still asking), 9404 Cancel -> only 9402
    r = P("/verify", {"grid": "urgent", "ids": [9404, 9401, 9403, 9402], "answers": ["no", "yes", "all_no", "cancel"]})
    assert r["verified"] == [9402] and r["messages"] == ["Əməliyyatlar verifikasiya edildi!"] and ids(r, "urgent") == [9401, 9403, 9404]
    # YesToAll + Yes: the rest without asking
    r = P("/verify", {"grid": "urgent", "ids": [9401, 9403, 9404], "answers": ["all"]})
    assert r["verified"] == [9401, 9403, 9404] and r["urgent"] == []


def test_verify_current_row_and_nothing():
    login(); P("/open")
    r = P("/verify", {"grid": "simple", "ids": [], "current": 9205, "answers": ["no"]})
    assert r["verified"] == [] and r["messages"] == [] and ids(r, "simple") == [9205]   # no call: no Refresh, no message
    r = P("/verify", {"grid": "simple", "ids": [], "current": 9205, "answers": ["yes"]})
    assert r["verified"] == [9205] and r["simple"] == []
    P("/verify", {"grid": "x", "ids": []}, 400)


def test_excel():
    login(); P("/open")
    x = client.post(B + "/excel", json={"grid": "simple"})
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.title == "Verification" and ws.max_row == 2
