"""Exchange Manager chunk 4 (Packages/ExchangeMGR-NBC) - api/app/em_turnover.py: the 'Dövriyyə' tab without its reports."""
import io
import pytest
from fastapi.testclient import TestClient
from app import main, sessions, em_turnover, em_sql as Q
from app.db import mock, mock_em

U = "/api/exchangemgr/turnover"
R = dict(beg="2026-10-01", end="2026-10-12")


@pytest.fixture(autouse=True)
def fresh():
    mock.reset(); mock_em.reset(); sessions.attempts._f.clear(); sessions.store._s.clear(); yield


def cl():
    c = TestClient(main.app); assert c.post("/api/login", json={"username": "ADMIN", "password": "Admin#2026"}).status_code == 200; return c


def post(c, url, code=200, **b):
    r = c.post(U + url, json=b); assert r.status_code == code, r.text; return r.json()


def test_lists():
    j = cl().get(U + "/lists").json()
    assert j["vdm_all"][0] == {"ex_no": -1, "ex_name": "Bütün məntəqələr"} and j["reestr_val"][0]["val_id"] == -1
    assert {v["valute"] for v in j["avalute"]} >= {1, 45}


def test_search_calls_get_turn_over_str_and_amountsum():
    c = cl()
    j = post(c, "/search", **R, val_id=None, status=0, ex_no=None)
    assert mock_em.CALLS[-1] == ("EXCH_MGR.GET_TURN_OVER_STR", ["2026-10-01", "2026-10-12", -1, 0, -1])
    assert len(j["rows"]) == 6 and j["rows"][0]["vdmname"] == "VDM-1 Nizami" and "custname" in j["rows"][0]
    assert len(j["sums"]) == 3 and j["sums"][0] == sum(r["debitamount"] for r in j["rows"])
    j = post(c, "/search", **R, val_id=1, status=1, ex_no=1)
    assert {(r["val_id"], r["status"], r["ex_id"]) for r in j["rows"]} == {(1, 1, 1)}
    assert j["sums"][0] == sum(r["debitamount"] for r in j["rows"])                  # S = St + Vt + Et into Amountsum.Sql[5]


def test_amountsum_line5_verbatim(monkeypatch):
    seen = []
    orig = mock.MockConnection.query
    monkeypatch.setattr(mock.MockConnection, "query", lambda self, sql, b=None: (seen.append(sql), orig(self, sql, b))[1])
    post(cl(), "/search", **R, val_id=45, status=2, ex_no=1)
    amt = [s for s in seen if s.startswith(Q.DM_Amountsum[:30])][0]
    assert amt.split("\n")[5] == "and (status=2)and (val_id=45) and (ex_id=1)"


def test_rate_turn_and_deleted():
    c = cl()
    j = post(c, "/rates", dat="2026-10-12")
    assert [r["tim"][11:16] for r in j["rows"] if r["val_id"] == 1] == ["08:30", "09:00"]   # ex_rates union ex_rates_archive
    j = post(c, "/deleted", **R, status=-1)
    assert mock_em.CALLS[-1][0] == "EXCH_MGR.GET_TURN_OVER_STRDEL" and [r["id"] for r in j["rows"]] == [900, 901]
    assert j["rows"][0]["status_str"] == "алыш" and j["rows"][0]["del_username"] == "ADMIN"


def test_toplam_and_reestr():
    c = cl()
    j = post(c, "/toplam", **R, val_on=True, val_id=1, val_name="USD", ex_on=True, ex_no=1)
    assert [r["ex_no"] for r in j["rows"]] == [1]
    r = j["rows"][0]
    assert r["val_name"] == "USD" and r["deltaamount"] == r["purchaseamount"] - r["salesamount"]
    assert len(post(c, "/toplam", **R)["rows"]) == 3                                   # unchecked -> -1 (unverified)
    assert post(c, "/reestr", dat="2026-10-12") == {"message": "em.err_choose_vdm_val"}
    assert post(c, "/reestr", dat="2026-10-12", val_id=1) == {"message": "em.err_choose_vdm"}
    assert post(c, "/reestr", dat="2026-10-12", ex_no=1) == {"message": "em.err_choose_currency"}
    j = post(c, "/reestr", dat="2026-10-12", ex_no=1, val_id=-1)
    assert "rows" in j


def test_excel_exports():
    from openpyxl import load_workbook
    c = cl()
    r = c.post(U + "/search/xlsx", json=dict(R, status=0))
    assert r.status_code == 200
    sh = load_workbook(io.BytesIO(r.content)).active
    assert sh.title == "Əməliyyatlar" and sh.max_row == 7 and "DEBITAMOUNT" in [x.value for x in sh[1]]
    assert load_workbook(io.BytesIO(c.post(U + "/rates/xlsx", json={"dat": "2026-10-12"}).content)).active.title == "Kurslar"
    assert load_workbook(io.BytesIO(c.post(U + "/deleted/xlsx", json=R).content)).active.title == "Silinmiş Əməliyyatlar"
    assert not [x for x in mock_em.CALLS if x[0] == "COMMIT"]
