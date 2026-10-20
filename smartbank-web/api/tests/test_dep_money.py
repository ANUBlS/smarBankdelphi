"""NewDepositsApplyMoney / NewDepositsExtractMoney + TrDelayed_Show - api/app/dep_money.py."""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from app import main, sessions
from app.db import mock, mock_dep


@pytest.fixture(autouse=True)
def fresh():
    mock.reset(); mock_dep.reset(); sessions.attempts._f.clear(); sessions.store._s.clear(); yield


def cl(u="ADMIN", p="Admin#2026"):
    c = TestClient(main.app); assert c.post("/api/login", json={"username": u, "password": p}).status_code == 200; return c


def post(c, url, **b):
    return c.post("/api/deposits" + url, json=b)


def test_apply_open_lists_and_accounts():
    o = post(cl(), "/apply/open", contract_id=5001).json()
    assert [t["id"] for t in o["tranches"]] == [71] and o["percent_text"] == "9"
    assert o["debit"]["id"] == 905001                                   # DebitAccs :custid := CustAccID (the client account)
    assert o["client_accs"][0] == {"id": -1, "acc_no": "Seçilməyib"} and 930001 in [a["id"] for a in o["client_accs"]]
    assert post(cl(), "/apply/open").json()["detail"]["key"] == "deposits.err_no_deposit"


def test_apply_ostat():
    c = cl()
    assert post(c, "/apply/ostat", accid=-1).json()["ostat"] == "0.00"
    assert post(c, "/apply/ostat", accid=930001).json()["ostat"] == "5400"


def test_apply_validations():
    c = cl()
    assert post(c, "/apply/run", contract_id=5001, total=0, percent_text="9").json()["detail"]["key"] == "deposits.m_err_disabled"
    assert post(c, "/apply/run", contract_id=5001, total=10, percent_text="0").json()["detail"]["key"] == "deposits.m_err_percent0"


def test_apply_confirm_commits_and_cancel_rolls_back():
    c = cl()
    r = post(c, "/apply/run", contract_id=5001, total=500, percent_text="9", descript="əmanət əlavə olunur", caccid=930001).json()
    assert r["caption"] == "Depozit yatrımı ..." and r["delayed"][0]["debitaccno"] == "38119440030001" and r["delayed"][0]["tramount"] == 500
    assert post(c, "/delayed/finish", ok=True).json()["committed"] is True
    assert [t["total"] for t in post(c, "/apply/open", contract_id=5001).json()["tranches"]] == [10000, 500]
    post(c, "/apply/run", contract_id=5001, total=700, percent_text="9", descript="x", caccid=-1)
    post(c, "/delayed/finish", ok=False)
    assert [t["total"] for t in post(c, "/apply/open", contract_id=5001).json()["tranches"]] == [10000, 500]
    assert post(c, "/delayed/finish", ok=True).status_code == 400          # nothing pending any more


def test_apply_delete_and_export_rights():
    c = cl()
    o = post(c, "/apply/delete", contract_id=5003, item_id=72).json()
    assert o["pending"] and [t["id"] for t in o["tranches"]] == [73]
    assert post(c, "/apply/delete", contract_id=5003, item_id=None).status_code == 400   # empty grid: DeleteDepositMoney(0)
    x = c.post("/api/deposits/apply/export.xlsx", json={"contract_id": 5003})
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    assert ws.title == "Depozit yatrımları" and [c.value for c in ws[1]] == ["Məbləğ", "Tarix", "Qeyd"]


def test_extract_add_defaults_and_withdraw():
    c = cl()
    a = post(c, "/extract/add", contract_id=5003).json()
    assert a["client"] == 3020.11 and a["main"] == 0 and a["date"] == "2026-10-12"
    assert a["custacc"]["contractaccno"] == "21019440005003"
    assert post(c, "/extract/run", contract_id=5003).json()["detail"]["key"] == "deposits.m_err_disabled"
    r = post(c, "/extract/run", contract_id=5003, main=1000, interest=100).json()
    assert [d["tramount"] for d in r["delayed"]] == [1000, 100]
    post(c, "/delayed/finish", ok=True)
    p = post(c, "/extract/open", contract_id=5003).json()
    assert p["payments"][-1]["facticaldebt"] == 1000 and p["info"]["facticaldebt"] == 49000
    assert "ORA-20103" in post(c, "/extract/run", contract_id=5003, main=10 ** 7).json()["detail"]["ora"]


def test_extract_delete_disabled_when_empty_and_export_titles():
    c = cl()
    assert post(c, "/extract/delete", contract_id=5001, item_id=0).json()["detail"]["key"] == "deposits.m_err_disabled"
    assert c.post("/api/deposits/extract/export.xlsx", json={"contract_id": 5001}).status_code == 400
    x = c.post("/api/deposits/extract/export.xlsx", json={"contract_id": 5003})
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    assert [c.value for c in ws[1]] == ["№", "Tarix", "Əsas məbləğdən", "Faiz məbləğindən", "CLIENTDEBT", "TRANCHESID", "Təyinat"]
    assert post(c, "/extract/delete", contract_id=5003, item_id=81).json()["payments"] == []


def test_stale_confirmation_is_dismissed_on_next_run():
    c = cl()
    post(c, "/apply/run", contract_id=5001, total=111, percent_text="9", caccid=-1)
    post(c, "/apply/run", contract_id=5001, total=222, percent_text="9", caccid=-1)
    post(c, "/delayed/finish", ok=True)
    assert [t["total"] for t in post(c, "/apply/open", contract_id=5001).json()["tranches"]] == [10000, 222]


def test_depositor_denied_money_rights_but_can_view():
    c = cl("DEPOSITOR", "Depositor#2026")
    assert post(c, "/apply/open", contract_id=5001).status_code == 200
    r = post(c, "/apply/add", contract_id=5001)
    assert r.status_code == 400 and "INSERT_DM_TRANSH" in r.json()["detail"]["ora"]
    assert "DELETE_DM_TRANSH" in post(c, "/right", code="DELETE_DM_TRANSH").json()["detail"]["ora"]
    assert post(c, "/right", code="DM_ADMIN").status_code == 404
