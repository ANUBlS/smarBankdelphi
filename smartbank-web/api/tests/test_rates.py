"""LastRates (OldRatesProj), RateDifference and RateDifferenceReport - api/app/oldrates.py, ratedifference.py,
ratedifferencereport.py."""
import io
import pytest
from fastapi.testclient import TestClient
from app import main, sessions, oldrates as O, ratedifferencereport as R, or_sql, rdr_sql
from app.db import mock
from app.delphi import subst


@pytest.fixture(autouse=True)
def fresh():
    mock.reset(); sessions.attempts._f.clear(); sessions.store._s.clear(); yield


def cl(u="ADMIN", p="Admin#2026"):
    c = TestClient(main.app); assert c.post("/api/login", json={"username": u, "password": p}).status_code == 200; return c


def post(c, url, code=200, **b):
    r = c.post(url, json=b); assert r.status_code == code, r.text; return r.json()


def test_sql_verbatim():
    assert or_sql.SQL_OR_ODSCOURSES.endswith("o.dat between :beg_dat and :end_dat\nOrder by :Title_Order")
    assert subst(or_sql.SQL_OR_ODSCOURSES, "Title_Order", "Val_kod").endswith("Order by Val_kod")
    assert rdr_sql.SQL_RDR_CURRENCY.startswith("Select  \n   *\nFrom\n   Avalute\nwhere \n  valute <> 0")
    s = R.refresh_sql(840, 101, 2)
    assert " and (Transactions.val_acc = 840 or Transactions.currencyID_c = 840) " in s
    assert " and (Debet.filial = 101 or Credit.filial = 101) " in s
    assert s.endswith(" and (Transactions.op_id = 2) Order by :Title ")
    assert R.refresh_sql(None, None, None) == R.REFRESH_HEAD + "Order by :Title "


def test_oldrates_init_and_period():
    c = cl()
    i = c.get("/api/oldrates/init").json()
    assert i["date"] == "2026-10-12" and i["title"] == "Val_kod"           # SBOP.Get_Date into both pickers
    assert [r["val_kod"] for r in i["rows"]] == sorted(r["val_kod"] for r in i["rows"]) and len(i["rows"]) == 5
    assert {r["dat"] for r in i["rows"]} == {"2026-10-12"}
    rows = post(c, "/api/oldrates/courses", beg="2026-10-06", end="2026-10-08", title="OSTAT")["rows"]
    assert len(rows) == 15 and [r["ostat"] for r in rows] == sorted(r["ostat"] for r in rows)
    assert all("2026-10-06" <= r["dat"] <= "2026-10-08" for r in rows)
    usd = next(r for r in rows if r["val_kod"] == 840); assert usd["val_name"] == "USD"
    # only grid FieldNames reach the substitution variable
    assert post(c, "/api/oldrates/courses", beg="2026-10-06", end="2026-10-12", title="1; drop table x")["title"] == "Val_kod"
    assert post(c, "/api/oldrates/courses", 422, beg="x", end="2026-10-12")


def test_oldrates_excel():
    from openpyxl import load_workbook
    r = cl().post("/api/oldrates/excel", json={"beg": "2026-10-12", "end": "2026-10-12"})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "Valyuta məzənnələri"
    assert [c.value for c in ws[1]] == ["Tarix", "$", "Valyutanın kodu", "Rəsmi məzənnə"]
    assert ws.max_row == 6


def test_ratedifference_never_opened():
    assert cl().get("/api/ratedifference/init").json() == {"active": False, "rows": []}


def test_rdr_init_and_refresh():
    c = cl()
    i = c.get("/api/ratedifferencereport/init").json()
    assert i["date"] == "2026-10-12" and i["title"] == "dat"
    assert 0 not in [x["id"] for x in i["currencies"]] and [x["id"] for x in i["currencies"]] == sorted(x["id"] for x in i["currencies"])
    assert i["filials"] and i["operators"]
    rows = post(c, "/api/ratedifferencereport/refresh", beg="2026-10-06", end="2026-10-12")["rows"]
    assert [r["id"] for r in rows] == [7001, 7002, 7003, 7004, 7005, 7006]      # 7007 has both amounts <> 0
    assert [r["recordno"] for r in rows] == list(range(1, 7))
    assert rows[0]["debetaccount"] == "AZ23SMRT10110000000000840101" and rows[0]["dcurrency"] == "USD" and rows[0]["ccurrency"] == "AZN"
    f = post(c, "/api/ratedifferencereport/refresh", beg="2026-10-06", end="2026-10-12", currency=840)["rows"]
    assert [r["id"] for r in f] == [7001, 7003, 7004, 7005]
    f = post(c, "/api/ratedifferencereport/refresh", beg="2026-10-06", end="2026-10-12", filial=102)["rows"]
    assert [r["id"] for r in f] == [7004]
    f = post(c, "/api/ratedifferencereport/refresh", beg="2026-10-06", end="2026-10-12", operator=2)["rows"]
    assert [r["id"] for r in f] == [7003, 7004]
    o = post(c, "/api/ratedifferencereport/refresh", beg="2026-10-06", end="2026-10-12", title="LAMMOUNT")
    assert o["title"] == "LAMMOUNT" and [r["lammount"] for r in o["rows"]] == sorted(r["lammount"] for r in o["rows"])
    assert post(c, "/api/ratedifferencereport/refresh", beg="2026-10-06", end="2026-10-12", title="RecordNo")["title"] == "dat"
    assert post(c, "/api/ratedifferencereport/refresh", 422, beg="2026-10-06", end="2026-10-12", currency="840 or 1=1")


def test_rights_denied():
    c = cl("TELLER", "Teller#2026")
    for u in ("/api/oldrates/init", "/api/ratedifference/init", "/api/ratedifferencereport/init"):
        assert c.get(u).status_code == 403
