"""BB payings (bbpayings.py, bb_sql.py, db/mock_bb.py)."""
import datetime
import pytest
from fastapi.testclient import TestClient
from app import main, sessions, bb_sql, bbpayings
from app.db import mock, mock_bb


@pytest.fixture(autouse=True)
def fresh():
    mock.reset(); sessions.attempts._f.clear(); sessions.store._s.clear(); yield


def cl(u="ADMIN", p="Admin#2026"):
    c = TestClient(main.app); assert c.post("/api/login", json={"username": u, "password": p}).status_code == 200; return c


def test_sql_verbatim():
    assert bb_sql.SQL_BB_ODS_ABNSERVICES.endswith("   and (sys_date>=(:StDate-1) and sys_date<=(:EndDate))")
    assert bbpayings.services_sql("AMOUNT").endswith("(:EndDate))\n order by AMOUNT")
    assert bb_sql.SQL_BB_ODS_SERVICES == "select * from BB_Services"


def test_init_window_uses_design_time():
    c = cl()
    i = c.get("/api/bbpayings/init").json()
    assert i["date"] == "2026-10-12" and len(i["contracts"]) == 3 and len(i["services"]) == 3
    # StDate-1 = 11.10 10:35:38 .. EndDate = 12.10 10:35:40 -> 104 (15:00) and 105 (05.10) are outside
    assert sorted(r["id"] for r in i["rows"]) == [101, 102, 103]
    o = c.post("/api/bbpayings/open", json={"st": "2026-10-05", "end": "2026-10-13", "order": "AMOUNT"}).json()["rows"]
    assert [r["amount"] for r in o] == sorted(r["amount"] for r in o) and len(o) == 5
    assert c.post("/api/bbpayings/open", json={"st": "2026-10-05", "end": "2026-10-13", "order": "x;--"}).status_code == 400


def test_add_update_delete_confirm():
    c = cl()
    n = c.post("/api/bbpayings/add", json={"contract_id": 13, "srv_id": 2, "amount": "7,5", "paymethod": 1}).json()["result"]
    assert mock_bb.CALLS[-1] == ("AbnServAdd", (13, None, 2, "7,5", 0, 0, datetime.datetime(2026, 10, 12), 1))
    c.post("/api/bbpayings/update", json={"id": n, "contract_id": 12, "srv_id": 1, "amount": "9", "paymethod": 0,
                                          "done_time": "2026-10-01T00:00:00"})
    assert mock_bb.CALLS[-1] == ("ABNSERVUPD", (n, 12, None, 1, "9", 0, 0, datetime.datetime(2026, 10, 1), 0))
    c.post("/api/bbpayings/delete", json={"id": 102, "answer": "cancel"})
    assert mock_bb.CALLS[-1][0] == "ABNSERVUPD"
    c.post("/api/bbpayings/delete", json={"id": 102, "answer": "no"})
    assert mock_bb.CALLS[-1] == ("ABNSERVDEL", (102, 0))
    c.post("/api/bbpayings/confirm", json={"ids": [103]})
    assert mock_bb.CALLS[-1] == ("ABNSERVCOMMIT", (103,))


def test_excel_filter():
    c = cl()
    x = c.post("/api/bbpayings/excel", json={"st": "2026-10-12", "end": "2026-10-12", "filter": {"status": 0}})
    from openpyxl import load_workbook
    import io
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.title == "Services" and ws.max_row == 3        # 102, 103


def test_rights_denied():
    c = cl("TELLER", "Teller#2026")
    assert c.get("/api/bbpayings/init").status_code == 403
    assert c.post("/api/bbpayings/confirm", json={"ids": [1]}).status_code == 403
