"""ASAN Finance (asanfinance.py, asan_sql.py, asan_http.py, db/mock_asan.py)."""
import io
import pytest
from fastapi.testclient import TestClient
from app import main, sessions, asan_sql
from app.db import mock, mock_asan


@pytest.fixture(autouse=True)
def fresh(monkeypatch):
    monkeypatch.setenv("ASAN_MODE", "mock")
    mock.reset(); sessions.attempts._f.clear(); sessions.store._s.clear(); yield


def cl(u="ADMIN", p="Admin#2026"):
    c = TestClient(main.app); assert c.post("/api/login", json={"username": u, "password": p}).status_code == 200; return c


def log():
    return mock_asan._S["log"]


def test_sql_verbatim():
    assert asan_sql.SQL_ASAN_IDDSET == "select \nseq_asan_loq.NEXTVAL id from dual"
    assert "and trunc(dat)>=trunc(sysdate)-3" in asan_sql.SQL_ASAN_LOGDSEST
    assert "JSON_TABLE(info, '$.Status[*]'" in asan_sql.SQL_ASAN_STATUS
    assert asan_sql.SQL_ASAN_ODS_LINK == "select*from asan_info"


def test_personal_fill_log_and_duplicate():
    c = cl()
    c.get("/api/asan/init")
    assert c.post("/api/asan/personal", json={"pin": "5ZX7K2P", "doc": ""}).json()["detail"]["key"] == "asan.err_pin_doc"
    j = c.post("/api/asan/personal", json={"pin": "5zx7k2p", "doc": "AA1234567"}).json()
    assert j["fill"]["PIN"] == "5ZX7K2P" and j["fill"]["Surname"] == "Həsənov" and j["fill"]["DocumentNumber"] == "AA1234567"
    assert j["image"] and j["conn_error"] is None
    last = log()[-1]
    assert last["STATUS"] == "success" and last["SERVICE"] == "Personalinfo" and last["DOCNUMBER"] == "AA1234567"
    assert last["URL"] == "http://prod.asanfinance.gov.az/api/v1/Info/Personal/5zx7k2p"        # BASE + url + Fin (no query)
    j2 = c.post("/api/asan/personal", json={"pin": "5zx7k2p", "doc": "AA1234567"}).json()        # cached within 3 days
    assert j2["fill"]["Surname"] == "Həsənov" and log()[-1]["STATUS"] == "duplicate" and log()[-1].get("DOCNUMBER") is None


def test_personal_status_and_connection_error():
    c = cl()
    j = c.post("/api/asan/personal", json={"pin": "BADPIN", "doc": "X1"}).json()
    assert j["message"] == "Fiziki şəxs tapılmadı"
    j = c.post("/api/asan/personal", json={"pin": "ERR500", "doc": "X1"}).json()
    assert j["conn_error"] == "HTTP/1.1 500 Internal Server Error"     # Memo1 unchanged -> the BADPIN answer again
    assert j["message"] == "Fiziki şəxs tapılmadı"


def test_balance_not_logged():
    c = cl()
    n = len(log()) if mock_asan._S else 1
    j = c.post("/api/asan/balance").json()
    assert j["balance"] == "1234.5" and len(log()) == n


def test_employee_pension_farm_voen():
    c = cl()
    e = c.post("/api/asan/employee", json={"pin": "5ZX7K2P"}).json()
    assert e["active"][0]["name"] == "Kənd MMC" and e["active"][0]["mqavile_bitme"] == "Müddətsiz" and e["deactive"][0]["enddate"] == "31.12.2019"
    p = c.post("/api/asan/pension", json={"pin": "7AB12CD"}).json()                 # log id 400 within 3 days -> duplicate
    assert p["pension"][0]["pendescription"] == "Əmək pensiyası (keşdən)" and log()[-1]["STATUS"] == "duplicate"
    f = c.post("/api/asan/farm", json={"pin": "5ZX7K2P", "doc": ""}).json()
    assert f["person"][0]["firstname"] == "Elçin" and [x["plname"] for x in f["farms"]] == ["Buğda", "Arpa"]
    v = c.post("/api/asan/voen", json={"pin": "1700054321"}).json()
    assert v["voen"][0]["voen"] == "1700054321" and v["voen"][0]["chief_surname"] == "Əliyev"
    assert c.post("/api/asan/employee", json={"pin": ""}).json()["detail"]["key"] == "asan.err_pin"


def test_expenses_and_excel():
    c = cl()
    j = c.post("/api/asan/expenses", json={"start": "2026-10-01 00:00:00", "end": "2026-10-12 10:00:00", "offset": 0, "limit": 10}).json()
    assert j["memo4"] == '{ "StartDate": "2026-10-01 00:00:00", "EndDate": "2026-10-12 10:00:00","Offset": 0,"Limit":10}'
    assert len(j["expenses"]) == 2 and log()[-1]["PIN"] == "post"
    x = c.post("/api/asan/expenses_excel", json={"id": j["id"]})
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.title == "Xercler" and ws.max_row == 3


def test_add_customer():
    c = cl()
    assert c.post("/api/asan/add_customer", json={"fields": {}}).json()["detail"]["key"] == "asan.err_query_first"
    f = {"PIN": "5ZX7K2P", "DocumentNumber": "AA1", "Name": "Elçin", "Surname": "Həsənov", "Patronymic": "Rauf oğlu",
         "BirthDate": "02.05.1985", "GivenDate": "15.03.2021", "ExpireDate": "15.03.2031", "Gender": "Qadın", "MaritalStatus": "Subay"}
    j = c.post("/api/asan/add_customer", json={"fields": f}).json()
    assert j["cust_id"] == 90001
    a = mock_asan.ADD_CUST[-1]
    assert a["p_sex"] == 2 and a["p_fam_status"] == 2 and a["p_passexpire"] == "15.03.2031" and a["p_birthday"].year == 1985
    f["BirthDate"] = "1985-05-02"
    assert c.post("/api/asan/add_customer", json={"fields": f}).json()["detail"]["key"] == "asan.err_date"


def test_rights_denied():
    c = cl("TELLER", "Teller#2026")
    assert c.get("/api/asan/init").status_code == 403
    assert c.post("/api/asan/balance").status_code == 403
