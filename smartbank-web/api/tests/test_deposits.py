import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from app import main, sessions, deposits as dp
from app.db import mock


@pytest.fixture(autouse=True)
def fresh():
    mock.reset(); sessions.attempts._f.clear(); sessions.store._s.clear(); yield


def cl(u="ADMIN", p="Admin#2026"):
    c = TestClient(main.app); assert c.post("/api/login", json={"username": u, "password": p}).status_code == 200; return c


def ids(c, **b):
    r = c.post("/api/deposits/list", json=b); assert r.status_code == 200, r.text; return [x["id"] for x in r.json()["rows"]]


# ------------------------------------------------------------------ BuildFilterString (NewDepositsManagerForm.pas:611-673)
def test_filter_string_empty_when_nothing_checked():
    assert dp.build_filter_string(dp.FilterIn()) == ""


def test_filter_string_currency_and_dates():
    s = dp.build_filter_string(dp.FilterIn(cb_currency=True, currency=944, cb_opendat=True, d1="2026-03-01", d2="2026-10-20"))
    assert s == "(CURRENCY=944) and (( opendate >=  '01.03.2026') and ( opendate <=  '20.10.2026'))"


def test_filter_string_category_maps_to_clienttype_and_person_to_personaltype():
    s = dp.build_filter_string(dp.FilterIn(cb_persons=True, category=2, cb_clienttype=True, person=1))
    assert "CLIENTTYPE=2" in s.upper() and "PERSONALTYPE=1" in s.upper()


def test_init_defaults_first_lookup_record_and_bank_date():
    i = cl().get("/api/deposits/init").json()
    assert i["defaults"]["currency"] == i["lookups"]["valutes"][0]["valute"]
    assert i["defaults"]["d1"] == i["bank_date"] == i["defaults"]["d2"]
    assert i["grid"] == dp.GRID


def test_list_unfiltered_shows_active_contracts_only():
    got = ids(cl())
    assert sorted(got) == list(range(5001, 5009)) and 4990 not in got
    assert got == sorted(got)                                     # order by cont.id


def test_filter_and_order():
    f = {"cb_currency": True, "currency": 944, "cb_opendat": True, "d1": "2026-03-01", "d2": "2026-10-20"}
    assert ids(cl(), filtered=True, filter=f, order="AMOUNT", desc=True) == [5003, 5008, 5005, 5007]


def test_filtered_false_ignores_checkboxes():
    assert len(ids(cl(), filtered=False, filter={"cb_currency": True, "currency": 840})) == 8


def test_order_whitelist():
    assert cl().post("/api/deposits/list", json={"order": "1;drop table x"}).status_code == 400


def test_status_filter_errors_like_delphi():
    r = cl().post("/api/deposits/list", json={"filtered": True, "filter": {"cb_status": True, "status_index": 0}})
    assert r.status_code == 400 and r.json()["detail"]["key"] == "deposits.err_filter"
    assert "STATUS" in r.json()["detail"]["ora"].upper()


def test_export_visible_fields_only():
    r = cl().post("/api/deposits/export.xlsx", json={})
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    head = [c.value for c in ws[1]]
    assert "DepozitinİD-si" not in head and "Müqavilənin №" in head
    assert ws.max_row == 9


def test_close_deposit_pending_and_removed_from_list():
    c = cl()
    r = c.post("/api/deposits/5002/close"); assert r.status_code == 200 and r.json()["pending"] is True
    assert 5002 not in ids(c)


def test_calc_toggle():
    c = cl()
    assert c.get("/api/deposits/5001/calc").json()["enabled"] is True
    assert c.post("/api/deposits/5001/calc", json={"enabled": True}).json()["pending"] is True
    assert c.get("/api/deposits/5001/calc").json()["enabled"] is False
    c.post("/api/deposits/5001/calc", json={"enabled": False})
    assert c.get("/api/deposits/5001/calc").json()["enabled"] is True


def test_plugins_not_migrated_and_unknown_key():
    c = cl()
    r = c.post("/api/deposits/plugin/contract_insert").json()
    assert r["migrated"] is False and r["bpl"] == "NewDepositsContract.bpl"
    assert c.post("/api/deposits/plugin/nope").status_code == 404


def test_rights_teller_denied_admin_ok():
    a = cl()
    assert a.post("/api/deposits/settings-menu").status_code == 200
    assert a.post("/api/deposits/plugin/sett_percent").status_code == 200
    t = cl("TELLER", "Teller#2026")
    assert t.get("/api/deposits/init").status_code == 403        # module not in TELLER's tree (require_module)
    assert t.post("/api/deposits/settings-menu").status_code == 403
