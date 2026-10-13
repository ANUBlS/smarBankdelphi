"""Packages/Out_Statement (TOut_StatementForm)."""
import io
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app
from app.db import mock as mock_db
from app.db import mock_ost as P
from app import out_statement as S
from app import out_ost_sql as Q

client = TestClient(app)
R = {"beg": "2026-10-06", "end": "2026-10-09"}


def login():
    mock_db.reset()
    assert client.post("/api/login", json={"username": "ADMIN", "password": "Admin#2026"}).status_code == 200


def test_update_str_and_flags_like_delphi():
    assert S.update_str("a, b,cc") == "'a','b','cc'"
    assert S.update_str("AZ1") == "'AZ1'" and S.update_str("  ") == "''" and S.update_str("") == "' '"
    # :894-901 with Pascal precedence: none checked -> ch_nothink, rate only -> ch_nothink
    assert [S.turn_and_ratediff(t, r) for t, r in ((True, True), (True, False), (False, True), (False, False))] == \
        ["ch_rate_and_turn", "ch_turn", "ch_nothink", "ch_nothink"]


def test_package_text_and_execution():
    login()
    j = client.post("/api/out-statement/run", json=R).json()
    assert j["turn_and_ratediff"] == "ch_turn" and j["bank"] == "SmartBank"
    accs = {a["accountnumber"][-6:]: a for a in j["accounts"]}
    assert accs["000022"]["rate_end"] == 1.7 and accs["000022"]["currencyid"] == 840
    assert all(t["operationcode"] != 1 for a in j["accounts"] for t in a["trans"])          # FilterTrans (OperationCode<>1)
    rd = client.post("/api/out-statement/run", json={**R, "rate_diff": True}).json()
    assert sum(len(a["trans"]) for a in rd["accounts"]) > sum(len(a["trans"]) for a in j["accounts"])
    # no off-balance movement on 12.10: notAccId lists every account -> all filtered when CBTurnOver is checked
    q = {"beg": "2026-10-12", "end": "2026-10-12"}
    assert client.post("/api/out-statement/run", json=q).json()["accounts"] == []
    off = client.post("/api/out-statement/run", json={**q, "turnover": False}).json()
    assert off["accounts"] and off["excluded"] == []
    usd = client.post("/api/out-statement/run", json={**R, "currency": 840}).json()
    assert {a["currencyid"] for a in usd["accounts"]} == {840}
    one = client.post("/api/out-statement/run", json={**R, "acc_no": "AZ34SMRT91018401010000000022"}).json()
    assert [a["accountid"] for a in one["accounts"]] == [50023]


def test_generated_sql_matches_package():
    import datetime
    d = datetime.datetime(2026, 10, 6)
    s = P.get_extraction(d, d, "-1", "'-1'", "-1", -1, -1, -1)
    assert "and ac.open_dat<='06-OCT-26'   order by ac.acc_no" in s
    s = P.get_extraction(d, d, "-1", "'AZ1'", "-1", 101, 840, 2)
    assert " and ac.val_acc=840 and ( ac.bal_cn1 in (-1) \n    or ac.acc_no in ('AZ1')" in s and s.endswith(" and ac.filial=101 and ac.OP_ID=2 order by ac.acc_no")
    t = P.extraction_trans(d, d, "-1", "'-1'", "-1", -1, -1, -1)
    assert "pr.debt=:OwnAccountId" in t and t.endswith(" order by dat,debit,kredit ")
    assert "nvl(vid_op,0)<>1" not in P.not_acc_id(d, d, "-1", "'-1'", "-1", -1, -1, -1, "ch_rate_and_turn")
    assert "nvl(vid_op,0)<>1" in P.not_acc_id(d, d, "-1", "'-1'", "-1", -1, -1, -1, "ch_turn")
    assert ":OwnAccountId" not in Q.SQL_OST_ODS_GETEXTRACTION and ":UserID" in Q.SQL_OST_ODS_EXTRACTIONTRANS


def test_validation_and_files():
    login()
    for bad in ({"acc_no": "x') or ('1'='1"}, {"bal_cn": "9101) or (1=1"}, {"cli_id": "1 or 1=1"}):
        assert client.post("/api/out-statement/run", json={**R, **bad}).status_code == 400
    assert client.post("/api/out-statement/statement.pdf", json=R).content[:4] == b"%PDF"
    assert client.post("/api/out-statement/statement.pdf", json={**R, "setting": 1, "print_mode": 1}).content[:4] == b"%PDF"
    txt = client.post("/api/out-statement/statement.txt", json=R).text
    assert "Hesab № AZ32SMRT91019441010000000020" in txt and "06.10.2026  09.10.2026" in txt
    x = client.post("/api/out-statement/statement.xlsx", json=R)
    ws = load_workbook(io.BytesIO(x.content)).active
    # first account AZN (t_val = 0): Cells[1,3] ' Hesab № :', [1,4] number, header row 7, first transaction row 8
    assert ws.cell(1, 3).value == " Hesab № :" and ws.cell(1, 4).value == "AZ32SMRT91019441010000000020"
    assert ws.cell(7, 2).value == "Tarix" and ws.cell(8, 5).value == 250.0
    assert ws.column_dimensions["G"].width == 100
