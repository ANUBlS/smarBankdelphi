"""Cards ▸ Yeni kartlar (cards_azericard U9.pas TForm9). newcard = CARD_CLIENT_APP messtype 1 and status 1 (fixture rows 9102, 9206)."""
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.db import mock as mock_db
from app.db import mock_cards as MC

client = TestClient(app)
B = "/api/cards/newcards"


def login():
    mock_db.reset(); MC.reset()
    assert client.post("/api/login", json={"username": "ADMIN", "password": "Admin#2026"}).status_code == 200


def P(path, body=None, code=200):
    r = client.post(B + path, json=body or {})
    assert r.status_code == code, r.text
    return r.json()


def test_open_lists_new_status1():
    login()
    r = P("/open")
    assert sorted(x["id"] for x in r["rows"]) == [9102, 9206]
    assert len(r["cols"]) == 25 and r["cols"][8] == "FIL" and r["cols"][21:] == ["", "", "", ""]
    assert r["titles_nofield"] == ["CREATED USER", "MODIFIED USER", "CREATE DATE", "MODIFY DATE"]


def test_delete_confirm_and_no():
    login(); P("/open")
    r = P("/delete", {"id": 9206, "confirmed": False})               # No: only the Refresh
    assert sorted(x["id"] for x in r["rows"]) == [9102, 9206]
    r = P("/delete", {"id": 9206, "confirmed": True})
    assert [x["id"] for x in r["rows"]] == [9102]
    P("/delete", {"id": 1, "confirmed": True}, 400)                    # not a row of the grid


def test_delete_before_open_and_excel():
    login()
    P("/delete", {"id": 9206, "confirmed": True}, 409)
    P("/open")
    x = client.post(B + "/excel")
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws.title == "CARDS" and "ROWID" not in [c.value for c in ws[1]] and ws.max_row == 3
