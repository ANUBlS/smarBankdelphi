"""Server-side PDF / Excel rendering of on-screen report grids (replaces QuickReport Preview and TExcelExporter).

Input is a neutral spec; all captions come from the caller (locale files on the web side), numbers are formatted with
2 decimals like the QRDBText DisplayFormat of the Delphi reports.
"""
from __future__ import annotations

import datetime
import io
from xml.sax.saxutils import escape
from decimal import Decimal

_FONT = None


def _font():
    global _FONT
    if _FONT is None:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        for name, path in (("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),):
            try:
                pdfmetrics.registerFont(TTFont(name, path))
                pdfmetrics.registerFont(TTFont(name + "-Bold", path.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")))
                _FONT = name
            except Exception:  # pragma: no cover - font missing in the image
                _FONT = "Helvetica"
    return _FONT


def fmt(v, money: bool) -> str:
    if v is None:
        return ""
    if money:
        return f"{Decimal(str(v)):,.2f}".replace(",", " ")
    if isinstance(v, datetime.datetime):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, datetime.date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _blocks(spec):
    """[(kind, payload)] rows to print: ('row', r) / ('gtotal', (key, totals)) / ('total', totals)."""
    out = []
    if spec.get("groups"):
        for g in spec["groups"]:
            out += [("row", r) for r in g["rows"]]
            out.append(("gtotal", (g["key"], g["totals"])))
    else:
        out += [("row", r) for r in spec["rows"]]
    if spec.get("totals"):
        out.append(("total", spec["totals"]))
    return out


def _line(spec, kind, payload, labels):
    cols = spec["cols"]
    if kind == "row":
        return [fmt(payload.get(c["field"]), c.get("money")) for c in cols]
    key, totals = (payload if kind == "gtotal" else (None, payload))
    cap = (labels.get("group_total", "") + " " + fmt(key, False)).strip() if kind == "gtotal" else labels.get("total", "")
    vals = [fmt(totals[c["field"]], True) if c["field"] in totals else "" for c in cols]
    vals[0] = cap
    return vals


def to_xlsx(spec: dict, labels: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    sheets = spec.get("sheets") or [spec]
    for n, sp in enumerate(sheets):
        ws = wb.active if n == 0 else wb.create_sheet()
        ws.title = (sp.get("sheet") or sp.get("title") or "Report")[:31]
        r = 1
        for h in [sp.get("title")] + list(sp.get("header") or []):
            if h:
                ws.cell(row=r, column=1, value=h).font = Font(bold=(r == 1), size=12 if r == 1 else 10)
                r += 1
        r += 0 if r == 1 else 1
        head = r
        for i, c in enumerate(sp["cols"], 1):
            cell = ws.cell(row=r, column=i, value=c["label"])
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0B2545")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.column_dimensions[cell.column_letter].width = max(12, min(48, len(str(c["label"])) + 4))
        for kind, payload in _blocks(sp):
            r += 1
            if kind == "row":
                for i, c in enumerate(sp["cols"], 1):
                    v = payload.get(c["field"])
                    if isinstance(v, Decimal):
                        v = float(v)
                    cell = ws.cell(row=r, column=i, value=v)
                    if c.get("money"):
                        cell.number_format = "#,##0.00"
            else:
                vals = _line(sp, kind, payload, labels)
                for i, (c, v) in enumerate(zip(sp["cols"], vals), 1):
                    totals = payload[1] if kind == "gtotal" else payload
                    val = float(totals[c["field"]]) if (i > 1 and c["field"] in totals) else (v if i == 1 else None)
                    cell = ws.cell(row=r, column=i, value=val)
                    cell.font = Font(bold=True)
                    if i > 1 and c.get("money"):
                        cell.number_format = "#,##0.00"
        ws.freeze_panes = ws.cell(row=head + 1, column=1)
        for t, nm in sp.get("signatures") or []:
            r += 2
            ws.cell(row=r, column=1, value=t)
            ws.cell(row=r, column=3, value=nm)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def to_pdf(spec: dict, labels: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    f = _font()
    wide = len(spec["cols"]) > 4
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=landscape(A4) if wide else A4, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=spec.get("title") or "")
    st_t = ParagraphStyle("t", fontName=f + "-Bold" if f != "Helvetica" else "Helvetica-Bold", fontSize=13, leading=16,
                          textColor=colors.HexColor("#0B2545"))
    st_h = ParagraphStyle("h", fontName=f, fontSize=9, leading=12)
    story = [Paragraph(escape(spec.get("title") or ""), st_t)]
    for h in spec.get("header") or []:
        story.append(Paragraph(escape(str(h)), st_h))
    story.append(Spacer(1, 4 * mm))
    data = [[c["label"] for c in spec["cols"]]]
    styles = [("FONTNAME", (0, 0), (-1, -1), f), ("FONTSIZE", (0, 0), (-1, -1), 8),
              ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2545")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
              ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9DB4C0")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for i, c in enumerate(spec["cols"]):
        if c.get("money"):
            styles.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
    for kind, payload in _blocks(spec):
        data.append(_line(spec, kind, payload, labels))
        if kind != "row":
            styles += [("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#E0F2F1")),
                       ("FONTNAME", (0, len(data) - 1), (-1, len(data) - 1), f + "-Bold" if f != "Helvetica" else "Helvetica-Bold")]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle(styles))
    story.append(t)
    for title, name in spec.get("signatures") or []:
        story += [Spacer(1, 6 * mm), Paragraph(f"{escape(title or '')} &nbsp;&nbsp; ____________________ &nbsp;&nbsp; {escape(name or '')}", st_h)]
    doc.build(story)
    return bio.getvalue()
