"""DepositsManager_New operation plugins with the delayed-transactions confirmation (SBDM_Common_new TrDelayed_Show):
  SBDM_ApplyMoney_new.bpl   (Yatırım)               A = frmApplyMoneynewImpl.pas, AM = frmAddMoneynewImpl.pas
  SBDM_ExtractMoney_new.bpl (Əmanətin ödənilməsi)   E = frmExtractMoneyImpl_new.pas, EM = frmAddExtractMoneyImpl_new.pas
  TT = SBDM_Common_new/TemporaryTransactionsFormImpl_new.pas, CD = SBDM_Common_new/CommonDeclsImpl_new.pas
SQL verbatim in dep_ops_sql.py. An operation runs in two HTTP calls on the same Oracle session, as the modal
TrDelayed_Show splits it in Delphi: (1) the package procedure fills SBCM_DELAYED_TRANSACTIONS_new and the rows are shown;
(2) Təsdiq et / İmtina et -> Commit<Op>(DoCommit => True/False); the package commits or rolls back itself.
"""
from __future__ import annotations
import datetime as dt
import io

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from pydantic import BaseModel, Field

from . import dep_ops_sql as Q
from . import deposits as D
from .customers import row
from .db import DBError

PKG_TR = "sbcm_credit_tranches_new"   # DMApplyMoneynewImpl.dfm:115-118 sbcm_tranches (pmNamed)
PKG_PAY = "sbcm_payment_new"          # DMExtractMoneyImpl_new.dfm:105-108 sbcm_payment (pmNamed)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# kind -> (package, commit procedure, TrDelayed_Show caption = rsTransactionsWillProcessed)
DELAYED = {"apply": (PKG_TR, "CommitDepositMoney", "Depozit yatrımı ..."),
           "extract": (PKG_PAY, "CommitWithDrawMoney", "Depozitin ödənilməsi ...")}
# CMExcelExport: visible TFields in field order (field, DisplayName, kind)
A_EXPORT = [("TOTAL", "Məbləğ", "f"), ("RESERVE", "Faiz", "f"), ("HDATE", "Tarix", "d"), ("DESCRIPTION", "Qeyd", "s")]
E_EXPORT = [("ID", "№", "f"), ("HDATE", "Tarix", "d"), ("FACTICALDEBT", "Əsas məbləğdən", "f"), ("INTERESTDEBT", "Faiz məbləğindən", "f")]
# ActnTransactionExecute AM:131-230 - result codes of the three check functions -> message (verbatim texts in the locales)
CHECK1 = {0: "a_chk1_0", 2: "a_chk1_2"}                                    # counttransh = 0: checkedtranch
CHECK2 = {0: "a_chk2_0", 1: "a_chk2_1", 2: "a_chk2_2", 4: "a_chk2_4", 6: "a_chk2_6", 7: "a_chk2_7"}   # counttransh > 0
CHECK3 = {0: "a_chk3_0", 1: "a_chk3_1", 2: "a_chk3_2", 3: "a_chk3_3", 4: "a_chk3_4"}                 # counttransh > 0


class OErr(Exception):
    def __init__(self, key, **kw):
        self.key, self.kw = key, kw


def _i(v):
    return 0 if v in (None, "") else int(float(v))


def bank_date(conn) -> dt.datetime:
    """Get_BankDate (the manager's BankDate dataset)."""
    v = conn.query(D.Q.SQL_BANKDATE)[0]["TEK_DAT"]
    return dt.datetime.combine(v.date() if isinstance(v, dt.datetime) else v, dt.time())


def selected(contract_id):
    if contract_id is None:
        raise OErr("deposits.err_no_deposit")          # Get_SelectedContID: 'Heç bir Depozit seçilməyib!'
    return int(contract_id)


def float_to_str(v) -> str:
    from .dep_contract import float_to_str as fs
    return "" if v is None else fs(v)


def xlsx(rows, cols, sheet) -> bytes:
    """CMExcelExport.ExportDataset(DataSet, SheetName, True, False): visible TFields, DisplayName titles."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = sheet[:31]
    for j, (_, lab, _) in enumerate(cols, 1):
        ws.cell(row=1, column=j, value=lab)
    for i, r in enumerate(rows, 2):
        for j, (f, _, k) in enumerate(cols, 1):
            v = r.get(f)
            if isinstance(v, dt.datetime):
                v = v.replace(tzinfo=None)
            c = ws.cell(row=i, column=j, value=v)
            c.number_format = "@" if k == "s" else "dd.mm.yyyy" if k == "d" else "#,##0.00" if f in ("TOTAL", "FACTICALDEBT", "INTERESTDEBT") else "General"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


# ------------------------------------------------------------------ Yatırım (SBDM_ApplyMoney_new)
def apply_open(conn, contract_id):
    """exportproc A:63-107: Tranches, DepositInfo, CustAcc (ClientACCID), DebitAccs(:CustID := CustAccID);
    frmAddMoneynew.AccountListCB.KeyValue := CustAccID. Caption = Get_SelectedContNo + '  nömrəli müqavilənin yatrımları'."""
    cid = selected(contract_id)
    tr = conn.query(Q.A_SQL_TRANCHES, {"ContractID": cid})
    info = conn.query(Q.A_SQL_DEPOSITINFO, {"ContractID": cid})
    ca = conn.query(Q.A_SQL_CUSTACC, {"Contractid": cid})
    cust_acc = _i(ca[0].get("CLIENTACCID")) if ca else 0
    debit = conn.query(Q.A_SQL_DEBITACCS, {"custid": cust_acc})
    return {"tranches": [row(x) for x in tr], "info": row(info[0]) if info else {}, "debit": [row(x) for x in debit],
            "accid": cust_acc}


def apply_add_open(conn, contract_id):
    """FormShow AM:77-92: DepositInfo.Refresh; PercentEdit := FloatToStr(getpersentTransh(ContractID));
    counttransh := counttranch(p_contractid); date := Get_BankDate; MoneySumEdit := 0."""
    cid = selected(contract_id)
    o = apply_open(conn, cid)
    pct = conn.call_func(PKG_TR, "getpersentTransh", [], float, named={"ContractID": cid})
    cnt = conn.call_func(PKG_TR, "counttranch", [], int, named={"p_contractid": cid})
    o.update(percent_text=float_to_str(pct if pct is not None else 0), counttransh=_i(cnt), date=bank_date(conn).date().isoformat(), sum=0)
    return o


class ApplyRun(BaseModel):
    contract_id: int | None = None
    accid: int | None = None              # AccountListCB.ListSource.DataSet ID (DebitAccs cursor)
    total: float = 0                      # MoneySumEdit.Value
    percent_text: str = Field("", max_length=64)
    counttransh: int = 0                  # read in FormShow


def apply_run(conn, b: ApplyRun):
    """ActnTransactionExecute AM:104-267 (ActnTransactionUpdate: Value > 0)."""
    if not b.total > 0:
        raise OErr("deposits.m_err_disabled")
    cid = selected(b.contract_id)
    try:
        pct = float(b.percent_text.strip().replace(",", "."))      # StrToFloat(PercentEdit.Text)
    except ValueError:
        raise OErr("deposits.c_err_number", value=b.percent_text)
    if pct == 0:
        raise OErr("deposits.m_err_percent0")                       # 'Yatrımın % dərəcəsi 0 ola bilməz.'
    nm = {"p_contractid": cid, "p_amuont": float(b.total)}
    c1 = _i(conn.call_func(PKG_TR, "checkedtranch", [], int, named=nm))
    c2 = _i(conn.call_func(PKG_TR, "checkedaddtranch", [], int, named=nm))
    c3 = _i(conn.call_func(PKG_TR, "checkedtranch_max_limit", [], int, named=nm))
    if b.counttransh == 0 and c1 in CHECK1:
        raise OErr("deposits." + CHECK1[c1])
    if b.counttransh > 0:
        for c, m in ((c2, CHECK2), (c3, CHECK3)):
            if c in m:
                raise OErr("deposits." + m[c])
    accid = _i(b.accid)
    conn.call_proc(PKG_TR, "depositmoney", [], named={
        "P_ClientAccID": accid, "contractid": cid, "hdate": bank_date(conn), "total": float(b.total), "PERCENT": pct,
        "Trname": "null", "usetransaction": True}, commit=False)


def apply_delete(conn, tranche_id):
    """ActnDelExecute A:270-287: DeleteDepositMoney(P_DepMoneyID => Tranches ID AsInteger), Refresh; no commit.
    ActnDel has no OnUpdate in the dfm (ActnDelUpdate is not wired): Sil is enabled on an empty grid (ID = 0)."""
    conn.call_proc(PKG_TR, "DeleteDepositMoney", [], named={"P_DepMoneyID": _i(tranche_id)}, commit=False)


# ------------------------------------------------------------------ Əmanətin ödənilməsi (SBDM_ExtractMoney_new)
def extract_open(conn, contract_id):
    """exportproc E: Payments, DepositInfo, CustAcc (ClientACCID), CreditAccs (InternalBankAccs zone; Locate ClientACCID ->
    AccountListCB.KeyValue). Caption = Get_SelectedContNo + ' nömrəli depozit ödənişləri'."""
    cid = selected(contract_id)
    pay = conn.query(Q.E_SQL_PAYMENTS, {"ContractID": cid})
    info = conn.query(Q.E_SQL_DEPOSITINFO, {"ContractID": cid})
    ca = conn.query(Q.E_SQL_CUSTACC, {"Contractid": cid})
    accs = [row(x) for x in conn.query(Q.E_SQL_CREDITACCS)]
    cust_acc = _i(ca[0].get("CLIENTACCID")) if ca else 0
    hit = any(_i(a.get("id")) == cust_acc for a in accs)
    return {"payments": [row(x) for x in pay], "info": row(info[0]) if info else {}, "custacc": row(ca[0]) if ca else {},
            "accs": accs, "accid": cust_acc if hit else None}


def extract_add_open(conn, contract_id):
    """FormShow EM: DepositInfo.Refresh; DTPOperation := Get_BankDate; EdMainDebit := 0; EdIntDebit := 0."""
    o = extract_open(conn, contract_id)
    o.update(date=bank_date(conn).date().isoformat(), main=0, interest=0)
    return o


class ExtractRun(BaseModel):
    contract_id: int | None = None
    accid: int | None = None              # AccountListCB.ListSource.DataSet ID (CreditAccs cursor)
    main: float = 0
    interest: float = 0


def extract_run(conn, b: ExtractRun):
    """ActiontrExecute EM (ActiontrUpdate: main + interest > 0): sbcm_payment_new.WithDrawMoney(..., UseTransactions,
    showTransaction => True)."""
    if not (b.main + b.interest) > 0:
        raise OErr("deposits.m_err_disabled")
    cid = selected(b.contract_id)
    conn.call_proc(PKG_PAY, "WithDrawMoney", [], named={
        "ContractID": cid, "P_AccID": _i(b.accid), "HDate": bank_date(conn), "FacticalDebtWithdraw": float(b.main),
        "InterestDebtWithdraw": float(b.interest), "UseTransactions": True, "showTransaction": True}, commit=False)


def extract_delete(conn, payment_id):
    """ActnDelExecute E: DeleteWithDrawMoney(P_DepWMoneyID => Payments ID), Refresh; no commit."""
    conn.call_proc(PKG_PAY, "DeleteWithDrawMoney", [], named={"P_DepWMoneyID": _i(payment_id)}, commit=False)


# ------------------------------------------------------------------ TrDelayed_Show (CD:76-87, TT)
def delayed_rows(conn):
    """FormShow TT:61-66: Transactions.Open; TransactionsCalcFields TT:93-100 TrAmount = DEBITAMOUNT else CREDITAMOUNT."""
    out = []
    for r in conn.query(Q.TT_SQL_TRANSACTIONS):
        x = row(r)
        x["tramount"] = x.get("debitamount") if r.get("DEBITAMOUNT") is not None else x.get("creditamount")
        out.append(x)
    return out


def delayed_finish(conn, kind: str, ok: bool):
    """Commit<Op>(DoCommit => TrDelayed_Show result); on an exception: Commit<Op>(False) and re-raise."""
    pkg, proc = DELAYED[kind][:2]
    try:
        conn.call_proc(pkg, proc, [], named={"DoCommit": bool(ok)}, commit=False)
    except DBError:
        if ok:
            conn.call_proc(pkg, proc, [], named={"DoCommit": False}, commit=False)
        raise


def start_delayed(ws, kind, fn, *a):
    """Operation procedure, then TrDelayed_Show's dataset. An error while showing -> Commit(False) (the Delphi except)."""
    prev = getattr(ws, "dep_delayed", None)
    if prev:                                   # a confirmation left open (page reload): it never committed - dismiss it
        ws.dep_delayed = None
        delayed_finish(ws.conn, prev, False)
    fn(ws.conn, *a)
    ws.dep_delayed = kind
    try:
        rows = delayed_rows(ws.conn)
    except DBError:
        ws.dep_delayed = None
        delayed_finish(ws.conn, kind, False)
        raise
    return {"delayed": rows, "caption": DELAYED[kind][2], "kind": kind}


# ------------------------------------------------------------------ router
class IdIn(BaseModel):
    contract_id: int | None = None
    item_id: int | None = None


class FinishIn(BaseModel):
    ok: bool


def make_router(current, require_module, err):
    r = APIRouter(prefix="/api/deposits")
    g = require_module(D.CLSID)

    def run(fn, *a):
        try:
            return fn(*a)
        except OErr as e:
            err(400, e.key, **e.kw)

    @r.post("/apply/open")
    def _a_open(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "NEW_SBCM_TRANSH")                        # miDepositClick (manager)
        return run(apply_open, ws.conn, b.contract_id)

    @r.post("/apply/add")
    def _a_add(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "INSERT_SBCM_TRANSH")                     # ActnAddExecute A:257-269
        return run(apply_add_open, ws.conn, b.contract_id)

    @r.post("/apply/run")
    def _a_run(b: ApplyRun, ws=Depends(g)):
        D.check_right(ws.conn, "INSERT_SBCM_TRANSH")
        return run(start_delayed, ws, "apply", apply_run, b)

    @r.post("/apply/delete")
    def _a_del(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "DELETE_SBCM_TRANSH")
        run(apply_delete, ws.conn, b.item_id)
        return {"ok": True, "pending": True, **run(apply_open, ws.conn, b.contract_id)}

    @r.post("/apply/export.xlsx")
    def _a_xlsx(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "EXCEL_SBCM_TRANSH")
        rows = ws.conn.query(Q.A_SQL_TRANCHES, {"ContractID": run(selected, b.contract_id)})
        return Response(xlsx(rows, A_EXPORT, "Depozit yatrımları"), media_type=XLSX,
                        headers={"Content-Disposition": 'attachment; filename="Depozit_yatrimlari.xlsx"'})

    @r.post("/extract/open")
    def _e_open(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "NEW_SBCM_PEYMENT")                       # miWithdrawClick (manager)
        return run(extract_open, ws.conn, b.contract_id)

    @r.post("/extract/add")
    def _e_add(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "INSERT_SBCM_PEYMENT")                    # ActnAddExecute E
        return run(extract_add_open, ws.conn, b.contract_id)

    @r.post("/extract/run")
    def _e_run(b: ExtractRun, ws=Depends(g)):
        D.check_right(ws.conn, "INSERT_SBCM_PEYMENT")
        return run(start_delayed, ws, "extract", extract_run, b)

    @r.post("/extract/delete")
    def _e_del(b: IdIn, ws=Depends(g)):
        if not run(extract_open, ws.conn, b.contract_id)["payments"]:
            err(400, "deposits.m_err_disabled")                          # ActnDelUpdate: RecordCount > 0
        D.check_right(ws.conn, "DELETE_SBCM_PEYMENT")
        run(extract_delete, ws.conn, b.item_id)
        return {"ok": True, "pending": True, **run(extract_open, ws.conn, b.contract_id)}

    @r.post("/extract/export.xlsx")
    def _e_xlsx(b: IdIn, ws=Depends(g)):
        rows = ws.conn.query(Q.E_SQL_PAYMENTS, {"ContractID": run(selected, b.contract_id)})
        if not rows:
            err(400, "deposits.m_err_disabled")                          # ActnExprt OnUpdate = ActnDelUpdate
        D.check_right(ws.conn, "EXCEL_SBCM_PEYMENT")
        return Response(xlsx(rows, E_EXPORT, "Ödənilmiş əməliyyat"), media_type=XLSX,
                        headers={"Content-Disposition": 'attachment; filename="Odenilmis_emeliyyat.xlsx"'})

    @r.post("/delayed/finish")
    def _finish(b: FinishIn, ws=Depends(g)):
        """Təsdiq et (ButtonCommitClick -> mrOk) or İmtina et / Esc after 'Əməliyyat aparılmayacaq! Əminsinizmi?'."""
        kind = getattr(ws, "dep_delayed", None)
        if not kind:
            err(400, "deposits.m_err_no_pending")
        ws.dep_delayed = None
        delayed_finish(ws.conn, kind, b.ok)
        return {"ok": True, "committed": b.ok, "kind": kind}

    return r
