"""Exchange & Convertation Reports - port of Delphi Packages/Exchange (TFrmExchange, Exchange.dpr; no other copies).

Units: UntExchangeImpl.pas (ActiveForm, 3 tabs), ExcangeDataModule.pas/.dfm (datasets; verbatim SQL in ex_sql.py),
AddExchangeUnit.pas (TFrmAddExchange 'VDM'), AddAccountUnit.pas (TFrmAddAccount 'Hesab'), UntTooling.pas (TFrmTool
'Exchange Balances', DBNavigator on ODSTooling), UntPrint.pas (TFrmPrint 'Report': Çap is on hold, its Excel button is
ported).

Not ported - QuickReport prints, on hold by the user (buttons stay and show the "on hold" notice):
  UntExchangeRep (ExchangeRep, TotalExchangeReport), QRepExchangeView (QRExchangeView, Exchange), UntQRConvertRep
  (QRConvertRep, ODSConvertBuy/Sell) and UntQRDetConvert (QRDetConvert, ODSConvertDetail) - FrmPrint 'Çap' and the
  Konvertasiya 'Alışlar' / 'Satışlar' buttons (ToolButtonBuyClick :853-878). Their SQL stays in ex_sql.py.

Package SB_Exchange (SQL_credit/sm_p.sql:42408-42518): add_exchanges and add_excaccount do NOT commit, upd_exchanges,
del_exchanges, upd_excaccount, del_excaccount and del_all_excaccount commit inside. Delphi never commits after them, so
every call here runs with commit=False (an added row stays in the user's open transaction, as in Delphi).
"""
from __future__ import annotations
import datetime
import io
import os
import re

from fastapi import APIRouter, Depends, Response
from pydantic import BaseModel, Field

from . import ex_sql as Q
from . import dbf_template as DBF
from .customers import row
from .db import DBError
from .delphi import date_only, delphi_round

CLSID = "{DC719027-B65B-11D6-A4B6-00E04C774264}"     # Exchange_TLB.pas:57 CLASS_FrmExchange
PKG = "SB_Exchange"                                   # ExcangeDataModule.dfm:5-7 OPacExchange.PackageName
NationalCurrency = "048"                              # UntExchangeImpl.pas:24
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXCHANGE_TITLE_DEFAULT = "Dat"                        # ExcangeDataModule.dfm:839-843 Exchange Variables.Data :TITLE = 'Dat'
CONVERT_TITLE_DEFAULT = "Assignment.Extra"            # ExcangeDataModule.dfm:189-193 ODSConvertation :TITLE
# Grid column FieldNames (UntExchangeImpl.dfm:333-486 ExchangeViewer, :658-813 ConvertationViewer) = Title values
EXCHANGE_COLS = ["DAT", "CURRENCYNAME", "BUYRATE", "BUYAMMOUNT", "BUYEQUIVALENT", "BUYCOUNT", "SELLRATE", "SELLAMMOUNT",
                 "SELLEQUIVALENT", "SELLCOUNT"]
CONVERT_COLS = ["DAT", "AMMOUNT", "AMMOUNT_C", "LAMMOUNT", "RATE_D", "MBANKKURS", "RATE_C", "DACC_NO", "KACC_NO", "FULLPATH",
                "DOFFICAL", "KCURRENCY", "MBANKKURS"]
# Persistent TFields (ExcangeDataModule.dfm) in field order: (FieldName, DisplayLabel, kind) - kind s=WideString,
# f=Float, i=Integer, d=DateTime. CMExcelExport.ExportDataset(DS, 'List1', True): VisibleFieldsOnly default True,
# DisplayName titles, column NumberFormat by DataType (CMExcelExport.pas:606-660).
F_EXCHANGE = [("DAT", "Tarix", "d"), ("CURRENCY", "Valyuta kodu", "f"), ("CURRENCYNAME", "Valyuta adı", "s"),
              ("INTERCODE", "Valyuta beynəlxalq kodu", "s"), ("BUYRATE", "Alış kursu", "f"), ("BUYAMMOUNT", "Alınmış valyuta", "f"),
              ("BUYEQUIVALENT", "Alınmış valyuta ekvivalent", "f"), ("MINBUYRATE", "Min alış kursu", "f"),
              ("MAXBUYRATE", "Maks alış kursu", "f"), ("BUYCOUNT", "Alış sayı", "f"), ("SELLRATE", "Satış kursu", "f"),
              ("SELLAMMOUNT", "Satılmış valyuta", "f"), ("SELLEQUIVALENT", "Satılmış valyuta ekvivalent", "f"),
              ("MINSELLRATE", "Min satış kursu", "f"), ("MAXSELLRATE", "Maks satış kursu", "f"), ("SELLCOUNT", "Satış sayı", "f")]
F_TOTAL = [("CURRENCY", "Valyuta kodu", "f"), ("CURRENCYNAME", "Valyuta adı", "s"), ("BUYCOUNT", "Alış sayı", "f"),
           ("SELLCOUNT", "Satış sayı", "f"), ("BUYAMMOUNT", "Alınmış valyuta", "f"), ("SELLAMMOUNT", "Satılmış valyuta", "f"),
           ("AVERAGEBUYRATE", "Orta alış kursu", "f"), ("AVERAGESELLRATE", "Orta satış kursu", "f"),
           ("MINBUYRATE", "Min alış kursu", "f"), ("MINSELLRATE", "Min satış kursu", "f"), ("MAXBUYRATE", "Max alış kursu", "f"),
           ("MAXSELLRATE", "Max satış kursu", "f")]
F_CONVERT = [("DAT", "Tarix", "d"), ("DOFFICAL", "Debet valyuta", "s"), ("AMMOUNT", "Debet məbləğ", "f"),
             ("KCURRENCY", "Kredit valyuta", "s"), ("AMMOUNT_C", "Kredit məbləğ", "f"), ("LAMMOUNT", "Ekvivalent", "f"),
             ("DACC_NO", "Debet hesab", "s"), ("KACC_NO", "Kredit hesab", "s"), ("RATE_D", "Sövdələşmə kursu", "f"),
             ("RATE_C", "Sövdələşmə kursu", "f"), ("MBANKKURS", "Mərkəzi Bankın kursu", "f"), ("FULLPATH", "Təyinat", "s"),
             ("VAL_ACC", "VAL_ACC", "f"), ("COFFICAL", "COFFICAL", "s"), ("SHORTNAME", "SHORTNAME", "s"), ("EXTRA", "EXTRA", "s")]
F_FULLEXCEL = [("SAY1", "Əqd", "f"), ("ORTA", "Orta", "s"), ("AVERAGERATE", "Orta kurs", "f"),
               ("AVERAGEAMOUNT", "Valyuta məbləği", "f"), ("EKV", "Ekvivalent", "f"), ("FULLPATH", "Konvertasiyanın təyinatı", "s"),
               ("VALNAME1", "Valyuta", "s")]      # FullexcelR1 'Sıra №' is Visible = False
# DOA default ApplyRecord of ODSTooling (UpdatingTable from 'From SetExchange SE'; ReadOnly False, CommitOnPost True)
SQL_TOOL_UPDATE = "update SetExchange set NAT_BALANCE = :NAT_BALANCE, FOR_BALANCE = :FOR_BALANCE where rowid = :ROWID_"
SQL_TOOL_INSERT = "insert into SetExchange (NAT_BALANCE, FOR_BALANCE) values (:NAT_BALANCE, :FOR_BALANCE)"
SQL_TOOL_DELETE = "delete from SetExchange where rowid = :ROWID_"
# Templates of the two BDE exports: ExtractFilePath(ParamStr(0)) = the SmartBank client folder in Delphi
EXCH_TEMPLATE = "EXCH122.dbf"                         # ActionFileExecute :1012
CONV_TEMPLATE = "$MB12212.dbf"                        # ToolButton24Click :756


def template_dir() -> str:
    from .config import settings
    d = os.getenv("DBF_TEMPLATE_DIR")
    if d:
        return d
    if settings.db_mode == "mock":
        return os.path.join(settings.mock_fixtures, "dbf")
    return os.path.join(os.path.dirname(__file__), "..", "templates")


class AccessIn(BaseModel):
    action: str = Field(..., pattern="^(Insert|Update|Delete)$")


class ExIn(BaseModel):
    id: int = 0                                          # ViewGrid.FieldByName('ID').AsInteger (0 on an empty grid)
    exnum: str = Field("", max_length=40)                # Edit1.Text 'VDM-in №'
    name: str = Field("", max_length=4000)               # Edit2.Text 'VDM-in adı'
    mode: int = Field(1, ge=1, le=2)                     # FrmExchange.i (1 = Yeni, 2 = Redaktə)


class IdIn(BaseModel):
    id: int = 0


class AccIn(BaseModel):
    exchange_id: int = 0                                 # ViewGrid ID
    acc_no: str = Field("", max_length=40)               # ComboBox2.Text
    mode: int = Field(3, ge=3, le=4)                     # FrmExchange.i (3 = Yeni, 4 = Redaktə)
    last_account_id: int = 0                             # Grid2.FieldByName('accountID').AsInteger


class AccDelIn(BaseModel):
    exchange_id: int = 0
    account_id: int = 0


class Period(BaseModel):
    beg: datetime.date | None = None                     # DateTimePicker1 / DTPCBeginDate (None = never set)
    end: datetime.date | None = None
    title: str | None = None                             # :Title (None = the dfm default)


class StatusIn(BaseModel):
    beg: datetime.date
    end: datetime.date
    status: str = Field(..., pattern="^(BUY|SELL)$")     # ToolButton26Click 'BUY' / ToolButton30Click 'SELL'


class FileIn(Period):
    filename: str = Field("", max_length=260)            # Edit1.Text / EConvert.Text (target file name)


class ToolIn(BaseModel):
    nat_balance: str = Field("", max_length=100)         # DBEdit1 NAT_BALANCE
    for_balance: str = Field("", max_length=100)         # DBEdit2 FOR_BALANCE


def strtoint(s: str) -> int:
    """SysUtils.StrToInt: optional sign, decimal or $hex; anything else -> EConvertError."""
    if not re.fullmatch(r"[-+]?(\d+|\$[0-9A-Fa-f]+)", s or ""):
        raise ValueError(f"'{s}' is not a valid integer value")
    neg = s.startswith("-")
    t = s.lstrip("+-")
    v = int(t[1:], 16) if t.startswith("$") else int(t)
    return -v if neg else v


def exchanges(c) -> dict:
    return {"exchanges": [row(x) for x in c.query(Q.SQL_EX_VIEWGRID)],     # ViewGrid
            "accounts": [row(x) for x in c.query(Q.SQL_EX_GRID2)]}          # Grid2 (filtered by ID in the page)


def exchange_rows(c, p: Period) -> list[dict]:
    """ToolButton13Click / ActionRefreshExecute :924-953 (Exchange.SetVariable BeginDate/EndDate, Open);
    ExchangeViewerTitleClick :905-922 (SetVariable('Title', FieldName); Refresh)."""
    return c.query(Q.SQL_EX_EXCHANGE, {"BeginDate": date_only(p.beg), "EndDate": date_only(p.end),
                                       "Title": EXCHANGE_TITLE_DEFAULT if p.title is None else p.title})


def convert_rows(c, p: Period) -> list[dict]:
    """ToolButton18Click :955-981 / ConvertationViewerTitleClick :711-728."""
    return c.query(Q.SQL_EX_ODSCONVERTATION, {"begindate": date_only(p.beg), "enddate": date_only(p.end),
                                              "title": CONVERT_TITLE_DEFAULT if p.title is None else p.title})


def xlsx(rows: list[dict], fields: list[tuple], sheet: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    for j, (_, lab, _k) in enumerate(fields, 1):
        ws.cell(row=1, column=j, value=lab).font = Font(name="Tahoma", bold=True)
    for i, r in enumerate(rows, 2):
        for j, (f, _, k) in enumerate(fields, 1):
            v = r.get(f)
            if isinstance(v, datetime.datetime):
                v = v.replace(tzinfo=None)
            cell = ws.cell(row=i, column=j, value=v)
            cell.number_format = {"s": "@", "i": "000000", "f": "0.00"}.get(k, "General")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _bank_code(c):
    r = c.query(Q.SQL_EX_ODSBANKCODE)                     # ODSBankCode.Open; FieldByName('NB').AsString
    v = r[0].get("NB") if r else None
    return "" if v is None else str(v)


def _template(name: str) -> bytes | None:
    p = os.path.join(template_dir(), name)
    try:
        with open(p, "rb") as fh:
            return fh.read()
    except OSError:
        return None


def get_foreign_currency(curr1: str, curr2: str) -> str:
    """GetForeignCurrency :737-743."""
    return curr2 if curr1 == NationalCurrency else curr1


def _as_string(v):
    """TField.AsString of the source dataset field (float -> FloatToStr, NULL -> '')."""
    if v is None:
        return ""
    if isinstance(v, float):
        from .delphi import float_to_str
        return float_to_str(v).replace(".", DBF.DECIMAL_SEPARATOR)
    return str(v)


def exchange_records(rows: list[dict], nb: str) -> list[dict]:
    """ActionFileExecute :1040-1068, one record per Exchange row."""
    out = []
    for e in rows:
        out.append({"WHEN": e.get("DAT"), "BANK_CODE": nb, "VALUTE": _as_string(e.get("INTERCODE")),
                    "RATE1": _as_string(e.get("SELLRATE")),
                    "TOTAL1": 0 if e.get("SELLAMMOUNT") is None else delphi_round(e["SELLAMMOUNT"]),
                    "MANAT1": _as_string(e.get("SELLEQUIVALENT")), "RATE2": _as_string(e.get("BUYRATE")),
                    "TOTAL2": 0 if e.get("BUYAMMOUNT") is None else delphi_round(e["BUYAMMOUNT"]),
                    "MANAT2": _as_string(e.get("BUYEQUIVALENT"))})
    return out


def convert_records(rows: list[dict], nb: str) -> list[dict]:
    """ToolButton24Click :791-841: one record (Tip 2) per row, a second one (Tip 1) when val_acc <> 0."""
    out = []
    for r in rows:
        if r.get("VAL_ACC") == 0:
            fcur, rate, amount = 1, r.get("RATE_C"), _as_string(r.get("AMMOUNT_C"))
        else:
            fcur, rate, amount = 0, r.get("RATE_D"), _as_string(r.get("AMMOUNT"))
        if rate is None:                                  # Real := Null variant -> EVariantTypeCastError
            raise ValueError("Could not convert variant of type (Null) into type (Double)")
        debt, credit = _as_string(r.get("COFFICAL")), _as_string(r.get("DOFFICAL"))
        base = {"BANK_CODE": nb, "INS1": get_foreign_currency(debt, credit), "INS2": NationalCurrency, "ALQ_SAT": fcur,
                "MABLAG": amount, "TARIX": r.get("DAT"), "KURS": float(rate)}
        out.append({"TIP": 2, "OPERATOR": _as_string(r.get("SHORTNAME")), **base})
        if r.get("VAL_ACC") != 0:
            out.append({"TIP": 1, "OPERATOR": _as_string(r.get("EXTRA")), **base})
    return out


def _file_name(fn: str, default: str) -> str:
    base = re.split(r"[\\/]", fn or "")[-1] or default
    return re.sub(r'[\r\n"]', "", base)


def make_router(current, require_module, err):
    r = APIRouter(prefix="/api/exchange")
    g = require_module(CLSID)

    def access(c, action):
        c.call_proc("SB_ACA", "CheckObjectAccess", ["Exchange", action], commit=False)

    @r.get("/init")
    def _init(ws=Depends(g)):
        """ActiveFormCreate :524-565."""
        c = ws.conn
        out = exchanges(c)                                              # ViewGrid.Active, Grid2.Active
        c.query(Q.SQL_EX_ODSTOOLING)                                    # ODSTooling.Open
        day = c.call_func("SB_MTOUTVIEW", "GetBankDay", [], datetime.datetime)   # getBankDay :489-496 (OPacCurrentBankDay)
        day = day.date() if isinstance(day, datetime.datetime) else day
        money = c.query(Q.SQL_EX_FORMONEY)                              # ForMoney -> Rate[] + FrmAddAccount.ComboBox1
        out.update(date=day.isoformat() if day else None,
                   currencies=[{"index": i, "valute": x.get("VALUTE"), "name": x.get("VALU_NAME")} for i, x in enumerate(money)],
                   exch_file=EXCH_TEMPLATE, conv_file=CONV_TEMPLATE, exchange_cols=EXCHANGE_COLS, convert_cols=CONVERT_COLS)
        return out

    @r.get("/grids")
    def _grids(ws=Depends(g)):
        return exchanges(ws.conn)

    @r.post("/access")
    def _access(b: AccessIn, ws=Depends(g)):
        """ToolButton6/10/8/1/5/3Click: SB_ACA.CheckObjectAccess('Exchange', ...) before the dialog / delete."""
        access(ws.conn, b.action)
        return {"ok": True}

    # ------------------------------------------------------------ TFrmAddExchange (Button2Click 'Qəbul')
    @r.post("/exchanges/save")
    def _ex_save(b: ExIn, ws=Depends(g)):
        c = ws.conn
        access(c, "Insert" if b.mode == 1 else "Update")
        if b.exnum == "" or b.name == "":
            err(400, "exchange.err_incomplete")                         # 'Вы не до конца ввели все данные!' + Abort
        try:
            n = strtoint(b.exnum)
        except ValueError as e:
            err(400, "exchange.err_convert", msg=str(e))
        if b.mode == 1:
            c.call_proc(PKG, "add_exchanges", [b.name, n], commit=False)
        else:
            c.call_proc(PKG, "upd_exchanges", [b.id, b.name, n], commit=False)
        return exchanges(c)                                             # ViewGrid.Refresh

    @r.post("/exchanges/delete")
    def _ex_delete(b: IdIn, ws=Depends(g)):
        """ToolButton8Click :587-604 (after 'VDM bütün hesabları ilə birlikdə silinsin?' = mrYes): del_all_excaccount
        only - the del_exchanges call is commented out in Delphi, so the VDM row itself stays."""
        c = ws.conn
        access(c, "Delete")
        c.call_proc(PKG, "del_all_excaccount", [b.id], commit=False)
        return exchanges(c)

    # ------------------------------------------------------------ TFrmAddAccount
    @r.get("/exaccounts/{index}")
    def _exacc(index: int, ws=Depends(g)):
        """ComboBox1Click: ExAccount.SetVariable('cod', Rate[ComboBox1.ItemIndex]) -> ComboBox2 items."""
        c = ws.conn
        money = c.query(Q.SQL_EX_FORMONEY)
        if not 0 <= index < len(money):
            err(400, "error.generic")
        cod = int(money[index]["VALUTE"])                               # Rate[] is array of Integer (AsInteger)
        return {"acc_no": [str(x["ACC_NO"]) for x in c.query(Q.SQL_EX_EXACCOUNT, {"cod": cod})]}

    @r.post("/exaccounts/save")
    def _acc_save(b: AccIn, ws=Depends(g)):
        """Button2Click: TakeAccountID(acc_cod := ComboBox2.Text); add_excaccount / upd_excaccount; Grid2.Refresh.
        Any exception -> 'The accounts are repeating!' #13 'One of the sectoes|sectores is free!' + Abort."""
        c = ws.conn
        access(c, "Insert" if b.mode == 3 else "Update")
        try:
            t = c.query(Q.SQL_EX_TAKEACCOUNTID, {"acc_cod": b.acc_no})
            acc_id = int(t[0]["ID"]) if t and t[0].get("ID") is not None else 0     # AsInteger on an empty dataset = 0
            if b.mode == 3:
                c.call_proc(PKG, "add_excaccount", [b.exchange_id, acc_id], commit=False)
            else:
                c.call_proc(PKG, "upd_excaccount", [b.exchange_id, acc_id, b.last_account_id], commit=False)
        except DBError as e:
            err(400, "exchange.err_repeating_3" if b.mode == 3 else "exchange.err_repeating_4", ora=e.message)
        return exchanges(c)

    @r.post("/exaccounts/delete")
    def _acc_delete(b: AccDelIn, ws=Depends(g)):
        """ToolButton3Click :606-616 - no confirmation."""
        c = ws.conn
        access(c, "Delete")
        c.call_proc(PKG, "del_excaccount", [b.exchange_id, b.account_id], commit=False)
        return exchanges(c)

    # ------------------------------------------------------------ V.D.M tab
    @r.post("/vdm/refresh")
    def _vdm(p: Period, ws=Depends(g)):
        return {"rows": [row(x) for x in exchange_rows(ws.conn, p)]}

    @r.post("/vdm/excel")
    def _vdm_xls(p: Period, ws=Depends(g)):
        """TFrmPrint 'Excel' with 'Ardıcıl hesabat' (ExportSource = ExchangeData: the open Exchange dataset)."""
        return Response(xlsx(exchange_rows(ws.conn, p), F_EXCHANGE, "Exchange"), media_type=XLSX,
                        headers={"Content-Disposition": 'attachment; filename="Exchange.xlsx"'})

    @r.post("/vdm/total-excel")
    def _vdm_total(p: Period, ws=Depends(g)):
        """TFrmPrint 'Excel' with 'Yekun hesabat': TotalExchangeReport re-opened with the variables ActionPrintExecute set
        from DateTimePicker1/2 (ExportSource = DataExcel)."""
        rows = ws.conn.query(Q.SQL_EX_TOTALEXCHANGEREPORT, {"BeginDate": date_only(p.beg), "EndDate": date_only(p.end)})
        return Response(xlsx(rows, F_TOTAL, "TotalExchange"), media_type=XLSX,
                        headers={"Content-Disposition": 'attachment; filename="TotalExchange.xlsx"'})

    @r.post("/vdm/file")
    def _vdm_file(b: FileIn, ws=Depends(g)):
        """ActionFileExecute :1002-1076. CopyFile(EXCH122.dbf) failing -> 'File Creation UnSuccesful Terminated!'."""
        c = ws.conn
        tpl = _template(EXCH_TEMPLATE)
        if tpl is None:
            err(400, "exchange.err_file_creation")
        rows = exchange_rows(c, b)
        nb = _bank_code(c)
        try:
            data = DBF.append(tpl, exchange_records(rows, nb))
        except (DBF.DbfError, ValueError) as e:
            err(400, "exchange.err_dbf", msg=str(e))
        return Response(data, media_type="application/x-dbf",
                        headers={"Content-Disposition": f'attachment; filename="{_file_name(b.filename, EXCH_TEMPLATE)}"'})

    # ------------------------------------------------------------ TFrmTool ('Sazlama')
    @r.get("/tooling")
    def _tool(ws=Depends(g)):
        rs = ws.conn.query(Q.SQL_EX_ODSTOOLING)
        return {"row": {"nat_balance": rs[0].get("NAT_BALANCE"), "for_balance": rs[0].get("FOR_BALANCE")} if rs else None}

    @r.post("/tooling")
    def _tool_post(b: ToolIn, ws=Depends(g)):
        """DBNavigator Edit (Insert when empty) + Post: DOA ApplyRecord, CommitOnPost."""
        c = ws.conn
        rs = c.query(Q.SQL_EX_ODSTOOLING)
        v = {"NAT_BALANCE": b.nat_balance or None, "FOR_BALANCE": b.for_balance or None}
        if rs:
            c.execute(SQL_TOOL_UPDATE, {**v, "ROWID_": rs[0]["ROWID"]})
        else:
            c.execute(SQL_TOOL_INSERT, v)
        c.commit()
        return _tool(ws)

    @r.post("/tooling/delete")
    def _tool_del(ws=Depends(g)):
        """DBNavigator Delete ('Delete record?' confirmed in the page): DOA delete + commit."""
        c = ws.conn
        rs = c.query(Q.SQL_EX_ODSTOOLING)
        if rs:
            c.execute(SQL_TOOL_DELETE, {"ROWID_": rs[0]["ROWID"]})
            c.commit()
        return _tool(ws)

    # ------------------------------------------------------------ Konvertasiya tab
    @r.post("/convert/refresh")
    def _conv(p: Period, ws=Depends(g)):
        return {"rows": [row(x) for x in convert_rows(ws.conn, p)]}

    @r.post("/convert/excel")
    def _conv_xls(p: Period, ws=Depends(g)):
        """ToolButton20Click :1116-1125: only when ODSConvertation is active and RecordCount > 0."""
        rows = convert_rows(ws.conn, p)
        if not rows:
            return Response(status_code=204)
        return Response(xlsx(rows, F_CONVERT, "List1"), media_type=XLSX,
                        headers={"Content-Disposition": 'attachment; filename="List1.xlsx"'})

    @r.post("/convert/excel-status")
    def _conv_status(b: StatusIn, ws=Depends(g)):
        """ToolButton26Click 'Excel alislar' / ToolButton30Click 'Excel satislar' :1127-1172 (Fullexcel)."""
        rows = ws.conn.query(Q.SQL_EX_FULLEXCEL, {"BEGINDATE": date_only(b.beg), "ENDDATE": date_only(b.end),
                                                  "status": b.status})
        if not rows:
            return Response(status_code=204)
        return Response(xlsx(rows, F_FULLEXCEL, "List1"), media_type=XLSX,
                        headers={"Content-Disposition": f'attachment; filename="List1_{b.status}.xlsx"'})

    @r.get("/convert/filename/{day}")
    def _fname(day: int, ws=Depends(g)):
        """DTPCEndDateChange :892-903: SaveConvert.FileName := GenerateFileName(Day)."""
        if not 1 <= day <= 31:
            err(400, "error.generic")
        return {"filename": ws.conn.call_func(PKG, "GenerateFileName", [day], str)}

    @r.post("/convert/file")
    def _conv_file(b: FileIn, ws=Depends(g)):
        """ToolButton24Click :745-851 (no tab check, no OnUpdate: also runs on an empty / never-opened dataset)."""
        c = ws.conn
        tpl = _template(CONV_TEMPLATE)
        if tpl is None:
            err(400, "exchange.err_file_creation")
        rows = convert_rows(c, b) if b.beg is not None else []      # never refreshed: variables NULL -> no rows
        nb = _bank_code(c)
        try:
            data = DBF.append(tpl, convert_records(rows, nb))
        except (DBF.DbfError, ValueError) as e:
            err(400, "exchange.err_dbf", msg=str(e))
        return Response(data, media_type="application/x-dbf",
                        headers={"Content-Disposition": f'attachment; filename="{_file_name(b.filename, CONV_TEMPLATE)}"'})

    return r
