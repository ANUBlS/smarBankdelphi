"""Cards ▸ Əməliyyat ▸ Excel yükləmələr sub forms (cards_azericard Unitexcel.pas TEXCELform Button1Click :45, by card_excelform.Id):
1 TSalary (Unitsalary.pas), 2 TNEWCARDEX (EXCELNEWCARD.pas), 3 TPROLONGATIONXLS (PROLONGATIONUNIT.pas), 4 Taccountclose (accountcloseunit.pas).

isertbtn 'Bazaya Göndər' of each form: OpenDialogCards.Execute, then Excel OLE: Workbooks.Open(file), Worksheets['CARDS'], one dataset
Insert per sheet row, Fields[i].AsString := Sheet.cells[j, cell] (Variant -> string = VarToStr), Post (DOA INSERT; CommitOnPost),
session.Commit after the loop. The web takes the ONE workbook uploaded by hand (.xlsx via openpyxl, .xls via xlrd; cached values,
not formulas) - choosing no file = the dialog cancelled. The DOA datasets select ROWID, which DOA does not expose as a field
(the QBEDefinition field lists of Unit18.dfm confirm it), so FieldCount counts the other select-list columns:
- salarybase  (Unit18.dfm:2438): CARDHOLDER, RBSNUMBER, AMMOUNT, CURRENCY, DEBET, STATUS, ID, INFO = 8 -> i := 0 to FieldCount-4 =
  CARDHOLDER..DEBET from columns 2..6, rows from 10 (Unitsalary.pas :86-101), until cells[j,2].Text = ''. No Trim. t := cells[6,3] (:105).
- YENICARD    (Unit18.dfm:2939): NAME..COMPANY, ID, STATUS, INFO = 14 -> i := 0 to 11 = NAME..COMPANY **and ID** from columns 2..13,
  rows from 2, Trim, until cells[j,2].Text = '' (EXCELNEWCARD.pas :93-108).
- PROLONG     (Unit18.dfm:2963): CARDNUMBER, DAT, STATUS, ID, INFO = 5 -> i := 0 to 2 = CARDNUMBER, DAT, **STATUS** from columns 1..3,
  rows from 2, Trim, until cells[j,1] and cells[j,2] are empty (PROLONGATIONUNIT.pas :82-100).
- acclose     (Unit18.dfm:3675): CARDNUMBER, STATUS, ID, INFO = 4 -> i := 0 to 1 = CARDNUMBER, **STATUS** from columns 1..2, same loop
  (accountcloseunit.pas :80-98).
The repeat .. until loop always inserts the first row (an empty sheet gives one empty record). A conversion error (TField.AsString on a
numeric / date field) stops the routine: the rows posted so far stay (CommitOnPost), nothing after the loop runs.
Refresh on a DataMod dataset that was never opened (PROLONG1, acclose2: no form opens them): the web assumes DOA re-executes the query
(as XÖHKS), so the grid fills. No right checks in the sub forms (Excel yükləmələr itself checks card_excelfile)."""
from __future__ import annotations

import base64
import datetime as dt
import io
import re

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from pydantic import BaseModel, Field

from . import cards_files as CF
from . import cards_sql as Q
from .cards import CLSID
from .cards_fileops import _dm_open
from .customers import row

DEC = CF.DEC

# form -> its two DataMod datasets, import layout and procedures (cited per form above / in FORMS comments)
FORMS = {
    # Unitsalary.pas: FormActivate :62 opens both; ToolButton1 :286 question -> transaction_salary; ToolButton2 :267 / FormClose :274
    # delete_salary; N11 / N12 ExportDataset(.., 'salary'); DBGrid1/2TitleClick :201/:222 (dgTitleClick + handler)
    "salary": dict(ds=(("salarybase", Q.LINES_SALARYBASE), ("salarybase2", Q.LINES_SALARYBASE2)), activate=True,
                   fields=["CARDHOLDER", "RBSNUMBER", "AMMOUNT", "CURRENCY", "DEBET"], row0=10, col0=2, stop=(2,), trim=False,
                   sheet="salary", sort=True,
                   cols=["CARDHOLDER", None, "RBSNUMBER", "AMMOUNT", "CURRENCY", "DEBET", "INFO"]),
    # EXCELNEWCARD.pas: FormActivate :293 opens both; FormShow :399 DateTimePicker1 := Date; ToolButton1 :286 newcommit; ToolButton2 :279
    # newdelet; N11 / N12 'YENICARD'; no OnTitleClick in the dfm (the handlers exist but are not wired)
    "newcard": dict(ds=(("YENICARD", Q.LINES_YENICARD), ("YENICARD1", Q.LINES_YENICARD1)), activate=True,
                    fields=["NAME", "SURNAME", "BIRTHDAY", "PASSPORT", "RESIDENT", "SECURITY", "PHONE", "CARDTYPE", "RBSNUMBER", "ADRESS",
                            "COMPANY", "ID"], row0=2, col0=2, stop=(2,), trim=True, sheet="YENICARD", sort=False,
                    cols=["ID", None, "NAME", "SURNAME", "BIRTHDAY", "PASSPORT", "RESIDENT", "SECURITY", "PHONE", "CARDTYPE", "RBSNUMBER",
                          "ADRESS", "COMPANY", "INFO"]),
    # PROLONGATIONUNIT.pas: no OnActivate; ToolButton1 :375 procommit; ToolButton2 :200 prodelet; N11 / N12 'CARDS'; no OnTitleClick
    "prolong": dict(ds=(("PROLONG", Q.LINES_PROLONG), ("PROLONG1", Q.LINES_PROLONG1)), activate=False,
                    fields=["CARDNUMBER", "DAT", "STATUS"], row0=2, col0=1, stop=(1, 2), trim=True, sheet="CARDS", sort=False,
                    cols=["ID", None, "CARDNUMBER", "DAT", "INFO"]),
    # accountcloseunit.pas: no OnActivate; ToolButton1 :119 acccommit; ToolButton2 :152 accdelet (isertbtn calls accdelete); 'CARDS'
    "accclose": dict(ds=(("acclose", Q.LINES_ACCLOSE), ("acclose2", Q.LINES_ACCLOSE2)), activate=False,
                     fields=["CARDNUMBER", "STATUS"], row0=2, col0=1, stop=(1, 2), trim=True, sheet="CARDS", sort=False,
                     cols=["ID", None, "CARDNUMBER", "INFO"]),
}
# TField types of the imported fields (the card_xls* DDL is not in the repo; the others are taken as strings)
FTYPE = {"AMMOUNT": "float", "ID": "float", "STATUS": "float", "DAT": "date"}
for _f in FORMS.values():                                  # DOA INSERT on these tables (mock_cards recognises the DML)
    m = re.search(r"\bfrom\s+(\w+)", Q.text(_f["ds"][0][1]), re.I)
    _f["table"] = m.group(1)
    CF.DML_TABLES.add(m.group(1).lower())
MAX_BYTES = 20 * 1024 * 1024


# ------------------------------------------------------------------------------------------------ Excel OLE replacement
class Sheet:
    """Worksheets['CARDS'] of the uploaded workbook; cells(row, col) 1-based -> the cell Value (None when empty)."""

    def __init__(self, data: bytes):
        self.cells = {}
        if data[:2] == b"PK":
            import openpyxl
            try:
                wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            except Exception:       # noqa: BLE001
                raise ValueError("file") from None
            ws = next((wb[n] for n in wb.sheetnames if n.upper() == "CARDS"), None)      # Excel sheet names: case-insensitive
            if ws is None:
                raise ValueError("Worksheets['CARDS']: Invalid index.")
            for r in ws.iter_rows():
                for c in r:
                    if getattr(c, "value", None) is not None:
                        self.cells[(c.row, c.column)] = c.value
            wb.close()
        elif data[:4] == b"\xd0\xcf\x11\xe0":
            import xlrd
            try:
                wb = xlrd.open_workbook(file_contents=data)
            except Exception:       # noqa: BLE001
                raise ValueError("file") from None
            ws = next((wb.sheet_by_name(n) for n in wb.sheet_names() if n.upper() == "CARDS"), None)
            if ws is None:
                raise ValueError("Worksheets['CARDS']: Invalid index.")
            for i in range(ws.nrows):
                for j in range(ws.ncols):
                    c = ws.cell(i, j)
                    if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    v = c.value
                    if c.ctype == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate.xldate_as_datetime(v, wb.datemode)
                    elif c.ctype == xlrd.XL_CELL_BOOLEAN:
                        v = bool(v)
                    elif c.ctype == xlrd.XL_CELL_ERROR:
                        v = None
                    self.cells[(i + 1, j + 1)] = v
        else:
            raise ValueError("file")

    def value(self, r, c):
        return self.cells.get((r, c))

    def text(self, r, c):
        return var_to_str(self.value(r, c))


def float_to_str(f: float) -> str:
    """FloatToStr (ffGeneral, 15 digits) with the client's DecimalSeparator."""
    s = "%.15G" % f
    if "E" in s:
        m, e = s.split("E")
        s = m + "E" + str(int(e))
    return s.replace(".", DEC)


def var_to_str(v) -> str:
    """Variant -> string of an Excel cell Value: empty '', varDouble FloatToStr, varDate DateToStr / DateTimeToStr (dd.mm.yyyy), varBoolean
    'True' / 'False', varOleStr as is."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, (int, float)):
        return float_to_str(float(v))
    if isinstance(v, dt.datetime):
        d = v.strftime("%d.%m.%Y")
        return d if v.time() == dt.time() else d + " " + f"{v.hour}:{v.minute:02d}:{v.second:02d}"
    if isinstance(v, dt.date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, dt.time):
        return f"{v.hour}:{v.minute:02d}:{v.second:02d}"
    return str(v)


def set_as_string(name: str, s: str):
    """TField.AsString := s for the field's type ('' -> Null)."""
    if s == "":
        return None
    t = FTYPE.get(name)
    if t == "float":
        try:
            if DEC != "." and "." in s:
                raise ValueError
            f = float(s.strip().replace(DEC, "."))
        except ValueError:
            raise ValueError(f"'{s}' is not a valid floating point value") from None
        return int(f) if f == int(f) else f
    if t == "date":
        m = re.fullmatch(r"\s*(\d{1,2})\.(\d{1,2})\.(\d{2,4})(?:\s+(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?)?\s*", s)
        try:
            y = int(m.group(3)); y = y + 2000 if y < 50 else y + 1900 if y < 100 else y
            return dt.datetime(y, int(m.group(2)), int(m.group(1)), int(m.group(4) or 0), int(m.group(5) or 0), int(m.group(6) or 0))
        except (AttributeError, ValueError):
            raise ValueError(f"'{s}' is not a valid date and time") from None
    return s


# ------------------------------------------------------------------------------------------------ datasets
def _st(ws, form):
    s = ws.module_state.get("cards_xls_" + form)
    if s is None:
        raise LookupError(form)
    return s


def _open(ws, form, k, sort=None):
    """Open / Refresh of DataMod dataset k (0 / 1) of the form; the open state outlives the form (DataMod)."""
    name, lines = FORMS[form]["ds"][k]
    L = list(lines) + (["ORDER BY " + sort] if sort else [])
    _st(ws, form)["rows"][k] = [dict(r) for r in ws.conn.query(Q.text(L), {})]
    _dm_open(ws).add(name.lower())


def _view(ws, form, **extra):
    s = _st(ws, form)
    out = {"grid1": None if s["rows"][0] is None else [row(r) for r in s["rows"][0]],
           "grid2": None if s["rows"][1] is None else [row(r) for r in s["rows"][1]], "cols": FORMS[form]["cols"],
           "sortable": FORMS[form]["sort"], "date": s.get("date")}
    out.update(extra)
    return out


def open_form(ws, form):
    if not ws.module_state.get("cards_excel"):
        raise LookupError(form)
    c = FORMS[form]
    ws.module_state["cards_xls_" + form] = s = {"rows": [None, None]}
    if form == "newcard":
        s["date"] = dt.date.today().isoformat()                   # FormShow :399 DateTimePicker1.DateTime := DATE
    for k, (name, _) in enumerate(c["ds"]):
        if c["activate"] or name.lower() in _dm_open(ws):          # FormActivate Open / still open from an earlier visit
            _open(ws, form, k)
    return _view(ws, form)


def _p(ws, name, args=()):
    ws.conn.call_proc("card_main", name, list(args), commit=False)          # DataMod_CARDS.mainpacket.CallProcedure


def _import(ws, form, sheet: Sheet) -> int:
    """The repeat .. until loop of isertbtnClick; returns j (the first row after the last imported one)."""
    c = FORMS[form]
    j = c["row0"]
    while True:
        rec = {}
        for i, f in enumerate(c["fields"]):
            v = sheet.text(j, c["col0"] + i)
            rec[f] = set_as_string(f, v.strip() if c["trim"] else v)
        cols = [k for k, v in rec.items() if v is not None]               # DOA INSERT of the assigned (non-Null) fields
        sql = f"insert into {c['table']} ({', '.join(cols)}) values ({', '.join(':b%d' % n for n in range(len(cols)))})" if cols else \
            f"insert into {c['table']} (STATUS) values (NULL)"
        ws.conn.execute(sql, {"b%d" % n: rec[k] for n, k in enumerate(cols)})
        ws.conn.commit()                                                    # Post (CommitOnPost)
        j += 1
        if all(sheet.text(j, col) == "" for col in c["stop"]):
            return j


def to_base(ws, form, data: bytes | None, date: str | None = None) -> dict:
    c, s = FORMS[form], _st(ws, form)
    msgs = []
    sheet = Sheet(data) if data is not None else None      # an unreadable file fails before anything is called (Workbooks.Open)
    try:
        if form == "newcard":                                   # EXCELNEWCARD.pas isertbtnClick :66
            d = dt.date.fromisoformat(date) if date else dt.date.today()
            s["date"] = d.isoformat()
            if not d > dt.date.today():                         # IF DateTimePicker1.DateTime > DATE (date part only: Kind dtkDate)
                return _view(ws, form, messages=["Tarixi düz seçin!"])
            _p(ws, "newdelet")
            if sheet is not None:
                _open(ws, form, 0)
                j = _import(ws, form, sheet)
                ws.conn.commit()
                _p(ws, "newupdate")
                _open(ws, form, 0); _open(ws, form, 1)
                msgs.append("%d record is imported!" % (j - 2))
            _p(ws, "updatedate", [dt.datetime(d.year, d.month, d.day)])
        elif form == "salary":                                  # Unitsalary.pas isertbtnClick :68
            t = ""
            if sheet is not None:
                _open(ws, form, 0)
                j = _import(ws, form, sheet)
                t = sheet.text(6, 3)
                ws.conn.commit()
                msgs.append("%d record is imported!" % (j - 10))
            _p(ws, "saltr", [t])
            _open(ws, form, 0); _open(ws, form, 1)
        else:                                                   # PROLONGATIONUNIT.pas :62 / accountcloseunit.pas :61
            pre, upd = ("prodelet", "proupdate") if form == "prolong" else ("accdelete", "accupdate")
            _p(ws, pre)
            if sheet is not None:
                _open(ws, form, 0)
                j = _import(ws, form, sheet)
                ws.conn.commit()
                _p(ws, upd)
                _open(ws, form, 0)
                if form == "prolong":
                    _open(ws, form, 0)                          # :111-112 PROLONG.Refresh twice (PROLONG1 is not refreshed)
                else:
                    _open(ws, form, 1)
                msgs.append("%d record is imported!" % (j - 2))
    except Exception as e:      # noqa: BLE001 - Delphi shows the exception; the rows posted so far stay
        return _view(ws, form, messages=msgs, error=getattr(e, "message", None) or str(e))
    return _view(ws, form, messages=msgs)


def action(ws, form, act, yes=False) -> dict:
    """ToolButton1 'Təsdiq Et' / ToolButton2 'Sil' / FormClose of the form, each followed by both datasets' Refresh."""
    _st(ws, form)
    procs = {("salary", "confirm"): "transaction_salary", ("salary", "delete"): "delete_salary", ("salary", "close"): "delete_salary",
             ("newcard", "confirm"): "newcommit", ("newcard", "delete"): "newdelet",
             ("prolong", "confirm"): "procommit", ("prolong", "delete"): "prodelet",
             ("accclose", "confirm"): "acccommit", ("accclose", "delete"): "accdelet"}
    p = procs.get((form, act))
    if p is None:
        raise LookupError(act)
    try:
        if not (form == "salary" and act == "confirm" and not yes):   # Unitsalary :288 MessageDlg .. = mrYes then transaction_salary
            _p(ws, p)
        _open(ws, form, 0); _open(ws, form, 1)
    except Exception as e:      # noqa: BLE001
        return _view(ws, form, error=getattr(e, "message", None) or str(e))
    if act == "close":
        ws.module_state.pop("cards_xls_" + form, None)
        return {"ok": True}
    return _view(ws, form)


def sort(ws, form, grid, field) -> dict:
    """TSalary DBGrid1/2TitleClick: IsEmpty or no FieldName -> Exit; SQL.Add('ORDER BY F'); Refresh; SQL[pos] := ''."""
    c, s = FORMS[form], _st(ws, form)
    k = 0 if grid == "grid1" else 1
    if not c["sort"] or not s["rows"][k] or field.upper() not in [x for x in c["cols"] if x]:
        return _view(ws, form)
    try:
        _open(ws, form, k, field.upper())
    except Exception as e:      # noqa: BLE001
        return _view(ws, form, error=getattr(e, "message", None) or str(e))
    return _view(ws, form)


def excel(ws, form, grid) -> bytes:
    """N11 / N12 (EXCEL1 / EXCEL2): CMExcelExport.ExportDataset(<dataset>, '<sheet>', True)."""
    from .cards_balpay import _excel
    c = FORMS[form]
    _st(ws, form)
    lines = c["ds"][0 if grid == "grid1" else 1][1]
    return _excel([dict(r) for r in ws.conn.query(Q.text(lines), {})], c["sheet"])


class UpIn(BaseModel):
    name: str = Field("", max_length=255)
    b64: str | None = Field(None, max_length=28_000_000)
    date: str | None = Field(None, max_length=10)


class ActIn(BaseModel):
    yes: bool = False


class SortIn(BaseModel):
    grid: str = Field(max_length=10)
    field: str = Field(max_length=40)


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_router(current, require_module, err):
    r = APIRouter(prefix="/api/cards/xls")
    g = require_module(CLSID)

    def _fe(fn, *a):
        try:
            return fn(*a)
        except LookupError:
            err(409, "error.bad_request")
        except ValueError as e:
            err(400, "error.bad_request", ora=str(e))

    def chk(form):
        if form not in FORMS:
            err(404, "error.bad_request")

    @r.post("/{form}/open")
    def _o(form: str, ws=Depends(g)):
        chk(form); return _fe(open_form, ws, form)

    @r.post("/{form}/tobase")
    def _b(form: str, b: UpIn, ws=Depends(g)):
        chk(form)
        data = None
        if b.b64 is not None:
            try:
                data = base64.b64decode(b.b64, validate=True)
            except Exception:       # noqa: BLE001
                err(400, "error.bad_request")
            if len(data) > MAX_BYTES:
                err(400, "error.bad_request")
        if b.date is not None and not re.fullmatch(r"\d{4}-\d\d-\d\d", b.date):
            err(400, "error.bad_request")
        return _fe(to_base, ws, form, data, b.date)

    for act in ("confirm", "delete", "close"):
        def mk(act):
            @r.post("/{form}/" + act)
            def _a(form: str, b: ActIn, ws=Depends(g)):
                chk(form); return _fe(action, ws, form, act, b.yes)
        mk(act)

    @r.post("/{form}/sort")
    def _s(form: str, b: SortIn, ws=Depends(g)):
        chk(form); return _fe(sort, ws, form, b.grid, b.field)

    @r.post("/{form}/excel")
    def _x(form: str, b: SortIn, ws=Depends(g)):
        chk(form)
        return Response(_fe(excel, ws, form, b.grid), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{FORMS[form]["sheet"]}.xlsx"'})

    return r
