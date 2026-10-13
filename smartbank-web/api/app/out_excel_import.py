"""Off-balance Excel import - port of Delphi Packages/OutExcelImport (TOutExcelImp, CLSID below).

OutExcelImporter.dpr compiles only OutExcelImport.pas. Package excelpkg (OraclePackage1, dfm). Flow:
- ActiveFormCreate (.pas:324-330): open OracleDataSet1, excelpkg.Del_outexcel, refresh.
- ToolButton1Click (.pas:502-597): Excel rows 2 .. SpecialCells(xlCellTypeLast).Row, columns 1..68 -> F1..F68, one dataset Insert per row.
  Excel OLE is replaced by openpyxl on the uploaded .xlsx (active sheet, values not formulas).
- ToolButton3Click (.pas:599-610): question, excelpkg.OUTEXCELTRANSACTIONS, message, Del_outexcel, refresh; SBSession.commit always.
- ToolButton5Click (.pas:613-620): reopen OracleDataSet1/OracleDataSet2.
- ToolButton7Click (.pas:622-628): excelpkg.del_outinspayments, message, Del_outexcel, refresh (no commit in the unit).
- DBGrid1 has dgEditing/dgConfirmDelete: cells are edited/deleted through the dataset (DOA rowid UPDATE/DELETE).
"""
from __future__ import annotations
import io
import re

from fastapi import APIRouter, Depends, Request
from pydantic import BaseModel, Field

from .customers import row

CLSID = "{9B81586F-881D-480A-AE1B-AAC7DF45622B}"     # OutExcelImport_TLB.pas:52 CLASS_OutExcelImp
PKG = "excelpkg"                                      # OutExcelImport.dfm OraclePackage1 PackageName
# OutExcelImport.dfm OracleDataSet1 SQL.Strings
SQL_LIST = 'select rowid,t. * from out_exceltable  t where userid=uid\r\n'
# OutExcelImport.dfm OracleDataSet2 SQL.Strings (opened by ToolButton5Click, not shown on the form)
SQL_VIDOP = 'select * from vid_op\r\n'
IMPORT_COLS = [f"F{i}" for i in range(1, 69)]                                   # .pas:519-586 f1..f68
GRID_COLS = IMPORT_COLS + ["USERID"] + [f"F{i}" for i in range(69, 76)]        # DBGrid1 Columns (dfm)
# DOA-generated statements for the dataset Insert / grid edit / grid delete (not in the source: DOA builds them at run time)
SQL_INSERT = "insert into out_exceltable (" + ", ".join(IMPORT_COLS) + ") values (" + ", ".join(":" + c for c in IMPORT_COLS) + ")"
SQL_DELETE = "delete from out_exceltable where rowid = :doa_rowid"
EDITABLE = set(IMPORT_COLS + [f"F{i}" for i in range(69, 76)])
MAX_BYTES = 20 * 1024 * 1024


def sql_update(col: str) -> str:
    if col not in EDITABLE:                         # column name is concatenated -> whitelist (F1..F75)
        raise ValueError("column")
    return f"update out_exceltable set {col} = :val where rowid = :doa_rowid"


def list_rows(conn) -> list[dict]:
    out = []
    for r in conn.query(SQL_LIST):
        d = row(r)
        d["rowid"] = r.get("ROWID")
        out.append(d)
    return out


def read_workbook(data: bytes) -> list[list]:
    """ToolButton1Click: m := 2; n := last used row + 1; while m<>n: Cells[m, 1..68].value."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValueError("file")
    ws = wb.active
    last = ws.max_row or 1
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=last, max_col=68, values_only=True):
        vals = list(r) + [None] * (68 - len(r))
        rows.append([None if (isinstance(v, str) and v == "") else v for v in vals[:68]])
    wb.close()
    return rows


def import_rows(conn, rows: list[list]) -> int:
    """One Insert/Post per row (CommitOnPost). An empty sheet makes Delphi's final Post fail ('not in edit or insert mode')."""
    if not rows:
        raise ValueError("empty")
    try:
        for vals in rows:
            conn.execute(SQL_INSERT, dict(zip(IMPORT_COLS, vals)))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return len(rows)


class EditIn(BaseModel):
    rowid: str = Field(max_length=64)
    col: str = Field(max_length=8)
    value: str | None = Field(default=None, max_length=4000)


def make_router(current, require_module, err):
    r = APIRouter(prefix="/api/out-excel-import")
    g = require_module(CLSID)

    @r.get("/init")
    def _init(ws=Depends(g)):
        c = ws.conn
        c.query(SQL_LIST)                                           # OracleDataSet1.Open
        c.call_proc(PKG, "Del_outexcel", [], commit=False)         # ActiveFormCreate .pas:327
        return {"cols": GRID_COLS, "rows": list_rows(c)}

    @r.get("/rows")
    def _rows(ws=Depends(g)):                                       # ToolButton5Click
        ws.conn.query(SQL_VIDOP)
        return {"rows": list_rows(ws.conn)}

    @r.post("/upload")
    async def _upload(request: Request, ws=Depends(g)):
        data = await request.body()
        if not data or len(data) > MAX_BYTES:
            err(400, "out_excel_import.err_file")
        try:
            n = import_rows(ws.conn, read_workbook(data))
        except ValueError as e:
            err(400, "out_excel_import.err_" + str(e))
        return {"imported": n, "rows": list_rows(ws.conn)}

    @r.post("/send")
    def _send(ws=Depends(g)):                                       # ToolButton3Click (after 'Yes')
        c = ws.conn
        c.call_proc(PKG, "OUTEXCELTRANSACTIONS", [], commit=False)
        c.call_proc(PKG, "Del_outexcel", [], commit=False)
        c.commit()                                                  # sboapiDataModule.SBSession.commit
        return {"rows": list_rows(c)}

    @r.post("/send-cancelled")
    def _send_no(ws=Depends(g)):                                    # ToolButton3Click 'No': only the commit runs
        ws.conn.commit()
        return {"rows": list_rows(ws.conn)}

    @r.post("/delete-payments")
    def _delpay(ws=Depends(g)):                                     # ToolButton7Click (no question in Delphi)
        c = ws.conn
        c.call_proc(PKG, "del_outinspayments", [], commit=False)
        c.call_proc(PKG, "Del_outexcel", [], commit=False)
        return {"rows": list_rows(c)}

    @r.post("/edit")
    def _edit(b: EditIn, ws=Depends(g)):                            # DBGrid1 dgEditing -> dataset Post (CommitOnPost)
        try:
            sql = sql_update(b.col.upper())
        except ValueError:
            err(400, "error.generic")
        c = ws.conn
        try:
            c.execute(sql, {"val": b.value, "doa_rowid": b.rowid})
            c.commit()
        except Exception:
            c.rollback()
            raise
        return {"rows": list_rows(c)}

    @r.post("/delete-row")
    def _delrow(b: dict, ws=Depends(g)):                            # DBGrid1 dgConfirmDelete (Ctrl+Del)
        rid = str(b.get("rowid") or "")
        if not re.fullmatch(r"[A-Za-z0-9+/]{1,64}", rid):
            err(400, "error.generic")
        c = ws.conn
        try:
            c.execute(SQL_DELETE, {"doa_rowid": rid})
            c.commit()
        except Exception:
            c.rollback()
            raise
        return {"rows": list_rows(c)}

    return r
