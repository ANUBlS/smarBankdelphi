"""NewDepositsApplyMoney.bpl (Yatırım) and NewDepositsExtractMoney.bpl (Əmanətin ödənilməsi), plus the shared
delayed-transactions confirmation (NewDepositsCommon TrDelayed_Show / TNewTransactionsViewForm).

Delphi units (NewDepositsManager/...):
  A  = NewDepositsApplyMoney/NewDepositsFrmApplyMoneyImpl.pas      AM = NewDepositsApplyMoney/NewDepositsFrmAddMoneyImpl.pas
  E  = NewDepositsExtractMoney/NewDepositsFrmExtractMoneyImpl.pas  EM = NewDepositsExtractMoney/NewDepositsFrmAddExtractMoneyImpl.pas
  TT = NewDepositsCommon/NewTemporaryTransactionsFormImpl.pas      CD = NewDepositsCommon/NewCommonDeclsImpl.pas
The operation runs in two HTTP steps on the same Oracle session, as the modal TrDelayed_Show splits it in Delphi:
  1. the package procedure writes DM_DELAYED_TRANSACTIONS (Usetransaction / showTransaction), the rows are shown;
  2. Təsdiq et / İmtina et -> Commit<Op>(DoCommit => True/False). The package does the commit / rollback itself.
"""
from __future__ import annotations
import datetime as dt
import io

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from pydantic import BaseModel

from . import dep_money_sql as Q
from . import deposits as D
from .customers import row
from .db import DBError

PKG_TR = "dm_credit_tranches"     # NewDepositsDMApplyMoneyImpl.dfm:95-100 dm_tranches, pmNamed
PKG_PAY = "dm_payment"            # NewDepositsDMExtractMoneyImpl.dfm:114-118 dm_payment, pmNamed
# kind -> (package, commit procedure, TrDelayed_Show caption)
DELAYED = {"apply": (PKG_TR, "CommitDepositMoney", "Depozit yatrımı ..."),           # AM rsTransactionsWillProcessed
           "extract": (PKG_PAY, "CommitWithDrawMoney", "Depozitin ödənilməsi ...")}   # EM rsTransactionsWillProcessed
DESCRIPTS = ["əmanət əlavə olunur", "əmanət əlavə olunur(nağd)"]                       # CBoxDescript Items.Strings
# Visible TFields in field order: (field, DisplayLabel, kind) - CMExcelExport ExportTitles
A_EXPORT = [("TOTAL", "Məbləğ", "f"), ("HDATE", "Tarix", "d"), ("DESCRIPTION", "Qeyd", "s")]
E_EXPORT = [("ID", "№", "f"), ("HDATE", "Tarix", "d"), ("FACTICALDEBT", "Əsas məbləğdən", "f"),
            ("INTERESTDEBT", "Faiz məbləğindən", "f"), ("CLIENTDEBT", "CLIENTDEBT", "f"), ("TRANCHESID", "TRANCHESID", "f"),
            ("DESCRIPTION", "Təyinat", "s")]


class MErr(Exception):
    def __init__(self, key, **kw):
        self.key, self.kw = key, kw


def _i(v):
    return 0 if v in (None, "") else int(float(v))


def bank_date(conn) -> dt.datetime:
    """Get_BankDate: the manager's BankDate dataset (select tek_dat from sysdat)."""
    r = conn.query(D.Q.SQL_BANKDATE)
    v = r[0]["TEK_DAT"]
    return dt.datetime.combine(v.date() if isinstance(v, dt.datetime) else v, dt.time())


def selected(contract_id):
    if contract_id is None:
        raise MErr("deposits.err_no_deposit")          # Get_SelectedContID: 'Heç bir Depozit seçilməyib!'
    return int(contract_id)


def float_text(v) -> str:
    """TDBEdit text of a TFloatField without DisplayFormat (FloatToStr, '.' decimal separator assumed)."""
    if v is None:
        return ""
    f = float(v)
    return str(int(f)) if f == int(f) else repr(f)


def xlsx(rows, cols, sheet) -> bytes:
    """CMExcelExport.ExportDataset(DataSet, SheetName, True, False): visible TFields, DisplayName titles."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = sheet[:31]
    for j, (_, lab, _) in enumerate(cols, 1):
        ws.cell(row=1, column=j, value=lab)
    for i, r in enumerate(rows, 2):
        for j, (f, _, k) in enumerate(cols, 1):
            v = r.get(f)
            if isinstance(v, dt.datetime):
                v = v.replace(tzinfo=None)
            c = ws.cell(row=i, column=j, value=v)
            c.number_format = "@" if k == "s" else "dd.mm.yyyy" if k == "d" else "General"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


# ------------------------------------------------------------------ Yatırım (ApplyMoney)
def apply_open(conn, contract_id):
    """exportproc A: Tranches, DepositInfo (CustID), CustAcc (ClientACCID), DebitAccs(:CustID := CustAccID), ClientAccs(:CustID := CustID)."""
    cid = selected(contract_id)
    tr = conn.query(Q.A_SQL_TRANCHES, {"ContractID": cid})
    info = conn.query(Q.A_SQL_DEPOSITINFO, {"ContractID": cid})
    cust_id = _i(info[0].get("CUSTOMERID")) if info else 0
    ca = conn.query(Q.A_SQL_CUSTACC, {"Contractid": cid})
    cust_acc = _i(ca[0].get("CLIENTACCID")) if ca else 0
    debit = conn.query(Q.A_SQL_DEBITACCS, {"custid": cust_acc})
    client = conn.query(Q.A_SQL_CLIENTACCS, {"custid": cust_id})
    i0 = info[0] if info else {}
    return {"tranches": [row(x) for x in tr], "info": row(i0) if i0 else {}, "percent_text": float_text(i0.get("INTEREST")),
            "debit": row(debit[0]) if debit else {}, "client_accs": [row(x) for x in client], "descripts": DESCRIPTS}


def apply_ostat(conn, accid):
    """DBLookupClientAccNoChange AM: KeyValue -1 -> '0.00', else AccNoOstat (ostat AsString)."""
    if accid is None or int(accid) == -1:
        return "0.00"
    r = conn.query(Q.A_SQL_ACCNOOSTAT, {"accid": int(accid)})
    v = r[0].get("OSTAT") if r else None
    return "" if v is None else float_text(v)


class ApplyRun(BaseModel):
    contract_id: int | None = None
    total: float = 0
    percent_text: str = ""
    descript: str = ""
    caccid: int | None = -1


def apply_run(conn, b: ApplyRun):
    """ActnTransactionExecute AM: ActnTransactionUpdate (sum > 0); PPercent = StrToFloat(DBEditPercent.Text); 0 -> message;
    the checkedtranch* checks are commented out; dm_credit_tranches.depositmoney(... 'Usetransaction', true)."""
    if not b.total > 0:
        raise MErr("deposits.m_err_disabled")
    cid = selected(b.contract_id)
    try:
        pct = float(b.percent_text.strip().replace(",", "."))
    except ValueError:
        raise MErr("deposits.c_err_number", value=b.percent_text)
    if pct == 0:
        raise MErr("deposits.m_err_percent0")         # 'Yatrımın % dərəcəsi 0 ola bilməz.'
    info = apply_open(conn, cid)
    conn.call_proc(PKG_TR, "depositmoney", [], named={
        "ClientAccID": _i(info["debit"].get("id")), "PAccID": -1 if b.caccid is None else int(b.caccid), "Contractid": cid,
        "Hdate": bank_date(conn), "Total": float(b.total), "Percent": pct, "Descript": b.descript, "Usetransaction": True},
        commit=False)


def apply_delete(conn, tranche_id):
    """ActnDelExecute A: right DELETE_DM_TRANSH, confirm, DeleteDepositMoney(P_DepMoneyID => Tranches ID), Refresh (no commit).
    ActnDelUpdate is not wired, so Sil is enabled on an empty grid too (ID AsInteger = 0)."""
    conn.call_proc(PKG_TR, "DeleteDepositMoney", [], named={"P_DepMoneyID": _i(tranche_id)}, commit=False)


# ------------------------------------------------------------------ Əmanətin ödənilməsi (ExtractMoney)
def extract_open(conn, contract_id):
    """exportproc E: Payments, DepositInfo, CustAcc, CreditAccs (opened, not shown: its combos are commented out)."""
    cid = selected(contract_id)
    pay = conn.query(Q.E_SQL_PAYMENTS, {"ContractID": cid})
    info = conn.query(Q.E_SQL_DEPOSITINFO, {"ContractID": cid})
    ca = conn.query(Q.E_SQL_CUSTACC, {"Contractid": cid})
    conn.query(Q.E_SQL_CREDITACCS)
    return {"payments": [row(x) for x in pay], "info": row(info[0]) if info else {}, "custacc": row(ca[0]) if ca else {}}


def extract_add_open(conn, contract_id):
    """FormShow EM: DepositInfo.Refresh; DTPOperation := Get_BankDate; EdClientDebit := SetResiduaryAmount(ContractID, Wdate)."""
    cid = selected(contract_id)
    o = extract_open(conn, cid)
    bd = bank_date(conn)
    fact = conn.call_func(PKG_PAY, "SetResiduaryAmount", [], float, named={"ContractID": cid, "Wdate": bd})
    o.update(date=bd.date().isoformat(), main=0, interest=0, client=fact)
    return o


class ExtractRun(BaseModel):
    contract_id: int | None = None
    main: float = 0
    interest: float = 0
    client: float = 0


def extract_run(conn, b: ExtractRun):
    """ActiontrExecute EM: ActiontrUpdate (main + int + client > 0); dm_payment.WithDrawMoney(..., 'showTransaction', 3)."""
    if not (b.main + b.interest + b.client) > 0:
        raise MErr("deposits.m_err_disabled")
    cid = selected(b.contract_id)
    conn.call_proc(PKG_PAY, "WithDrawMoney", [], named={
        "ContractID": cid, "HDate": bank_date(conn), "FacticalDebtWithdraw": float(b.main),
        "InterestDebtWithdraw": float(b.interest), "ClientDebtWithdraw": float(b.client), "showTransaction": 3}, commit=False)


def extract_delete(conn, payment_id):
    """ActnDelExecute E: right DELETE_SBCM_PEYMENT, confirm, DeleteWithDrawMoney(P_DepWMoneyID), Refresh (no commit)."""
    conn.call_proc(PKG_PAY, "DeleteWithDrawMoney", [], named={"P_DepWMoneyID": _i(payment_id)}, commit=False)


# ------------------------------------------------------------------ TrDelayed_Show
def delayed_rows(conn):
    """TNewTransactionsViewForm.FormShow: Transactions.Open; TransactionsCalcFields TrAmount = DEBITAMOUNT else CREDITAMOUNT."""
    out = []
    for r in conn.query(Q.SQL_DELAYED):
        x = row(r)
        x["tramount"] = x.get("debitamount") if r.get("DEBITAMOUNT") is not None else x.get("creditamount")
        out.append(x)
    return out


def delayed_finish(conn, kind: str, ok: bool):
    """Commit<Op>(DoCommit => TrDelayed_Show result); on an exception: Commit<Op>(False) and re-raise."""
    pkg, proc, _ = DELAYED[kind]
    try:
        conn.call_proc(pkg, proc, [], named={"DoCommit": bool(ok)}, commit=False)
    except DBError:
        if ok:
            conn.call_proc(pkg, proc, [], named={"DoCommit": False}, commit=False)
        raise


# ------------------------------------------------------------------ router
class IdIn(BaseModel):
    contract_id: int | None = None
    item_id: int | None = None
    accid: int | None = None


class FinishIn(BaseModel):
    ok: bool


def make_router(current, require_module, err):
    r = APIRouter(prefix="/api/deposits")
    g = require_module(D.CLSID)

    def run(fn, *a):
        try:
            return fn(*a)
        except MErr as e:
            err(400, e.key, **e.kw)

    def start(ws, kind, fn, b):
        """Step 1: the operation procedure, then TrDelayed_Show's dataset (inside Delphi's try: an error -> Commit(False))."""
        prev = getattr(ws, "dep_delayed", None)
        if prev:                                       # a confirmation left open (page reload): it never committed - dismiss it
            ws.dep_delayed = None
            delayed_finish(ws.conn, prev, False)
        run(fn, ws.conn, b)
        ws.dep_delayed = kind
        try:
            rows = delayed_rows(ws.conn)
        except DBError:
            ws.dep_delayed = None
            delayed_finish(ws.conn, kind, False)
            raise
        return {"delayed": rows, "caption": DELAYED[kind][2]}

    @r.post("/apply/open")
    def _a_open(b: IdIn, ws=Depends(g)):
        return run(apply_open, ws.conn, b.contract_id)

    @r.post("/apply/add")
    def _a_add(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "INSERT_DM_TRANSH")                     # ActnAddExecute A
        o = run(apply_open, ws.conn, b.contract_id)                    # FormShow AM: DepositInfo.Refresh
        return {**o, "sum": "0.00", "descript_index": 0, "caccid": -1, "ostat": "0.00", "date": bank_date(ws.conn).date().isoformat()}

    @r.post("/apply/ostat")
    def _a_ostat(b: IdIn, ws=Depends(g)):
        return {"ostat": apply_ostat(ws.conn, b.accid)}

    @r.post("/apply/run")
    def _a_run(b: ApplyRun, ws=Depends(g)):
        return start(ws, "apply", apply_run, b)

    @r.post("/apply/delete")
    def _a_del(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "DELETE_DM_TRANSH")
        apply_delete(ws.conn, b.item_id)
        return {"ok": True, "pending": True, **run(apply_open, ws.conn, b.contract_id)}

    @r.post("/apply/export.xlsx")
    def _a_xlsx(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "EXCEL_DM_TRANSH")
        cid = run(selected, b.contract_id)
        rows = ws.conn.query(Q.A_SQL_TRANCHES, {"ContractID": cid})
        return Response(xlsx(rows, A_EXPORT, "Depozit yatrımları"), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="Depozit_yatrimlari.xlsx"'})

    @r.post("/extract/open")
    def _e_open(b: IdIn, ws=Depends(g)):
        return run(extract_open, ws.conn, b.contract_id)

    @r.post("/extract/add")
    def _e_add(b: IdIn, ws=Depends(g)):
        D.check_right(ws.conn, "INSERT_SBCM_PEYMENT")                  # ActnAddExecute E
        return run(extract_add_open, ws.conn, b.contract_id)

    @r.post("/extract/run")
    def _e_run(b: ExtractRun, ws=Depends(g)):
        return start(ws, "extract", extract_run, b)

    @r.post("/extract/delete")
    def _e_del(b: IdIn, ws=Depends(g)):
        o = run(extract_open, ws.conn, b.contract_id)
        if not o["payments"]:
            err(400, "deposits.m_err_disabled")                         # ActnDelUpdate: RecordCount > 0
        D.check_right(ws.conn, "DELETE_SBCM_PEYMENT")
        extract_delete(ws.conn, b.item_id)
        return {"ok": True, "pending": True, **run(extract_open, ws.conn, b.contract_id)}

    @r.post("/extract/export.xlsx")
    def _e_xlsx(b: IdIn, ws=Depends(g)):
        cid = run(selected, b.contract_id)
        rows = ws.conn.query(Q.E_SQL_PAYMENTS, {"ContractID": cid})
        if not rows:
            err(400, "deposits.m_err_disabled")                         # ActnExprt OnUpdate = ActnDelUpdate
        D.check_right(ws.conn, "EXCEL_SBCM_PEYMENT")
        return Response(xlsx(rows, E_EXPORT, "Ödənilmiş əməliyyat"), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="Odenilmis_emeliyyat.xlsx"'})

    @r.post("/delayed/finish")
    def _finish(b: FinishIn, ws=Depends(g)):
        """Təsdiq et (ButtonCommitClick -> mrOk) or İmtina et / Esc after 'Əməliyyat aparılmayacaq! Əminsinizmi?' (FormCloseQuery)."""
        kind = getattr(ws, "dep_delayed", None)
        if not kind:
            err(400, "deposits.m_err_no_pending")
        ws.dep_delayed = None
        delayed_finish(ws.conn, kind, b.ok)
        return {"ok": True, "committed": b.ok, "kind": kind}

    return r
