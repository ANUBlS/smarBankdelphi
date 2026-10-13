"""Off-balance account statement - port of Delphi Packages/Out_Statement/Ocx (TOut_StatementForm, CLSID below).

Form Out_StatementImpl.pas, data module Out_Stdatamod_Unit.pas (SQL: out_ost_sql.py, verbatim from Out_Stdatamod_Unit.dfm),
reports Out_RepR_Unit (TFormRepR, purpose on the right) / Out_Rep_Unit (TFormRep, purpose below), Out_FormTransView_Unit.
ExportDataTable :879-1066 is reproduced step by step in export_data_table():
  * Turn_and_RateDiff :894-901 (with Delphi operator precedence, see turn_and_ratediff)
  * UpdateStr :382-431 quoting of the account / balance lists, '-1' / '''-1''' defaults :914-928
  * OUT_SB_STATEMENT.notAccId (:939) returns SQL text; when CBTurnOver is checked that text is opened (ODSNotAccId) and
    every accountid becomes '(accountid<>N) and ' of the ODSMainState client filter (:956-971)
  * ODS_GetExtraction / ODS_ExtractionTrans (:992-1029) return a STR column with DB-generated SQL; it is executed as
    returned (ODSMainState; ODSExtractTrans per master row with :OwnAccountId = AccountId, MasterFields/DetailFields)
  * FilterTrans '(OperationCode<>1)' when CBRateDiff is unchecked (:932-935), applied client side like TDataSet.Filter
  * RatesDataSet / RatesDataSetBegin: Operkurs for (CurrencyID, EndDate) / (CurrencyID, DatDate) (master-detail links)
The text the user typed is bound (the package concatenates it), but is validated first: account numbers [A-Za-z0-9],
balance accounts / client ids integers - for valid input the generated SQL is identical to Delphi's.
Outputs (BtnOkClick :1068-1080, CBoxExtractCheck): 0 preview (RGroupSetting 0 Out_RepR / 1 Out_Rep) -> PDF,
1 'fayla yaz' (QRAsciiExportFilter of Out_Rep) -> .txt, 2 Excel OLE (:519-833) -> openpyxl with the same cell layout,
3 transactions view (FormTransView), 4 'plastik kart' -> ShowCardState just closes the form.
"""
from __future__ import annotations
import datetime
import io
import re
from decimal import Decimal

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from pydantic import BaseModel, Field

from . import out_ost_sql as Q
from . import report_render as R
from .balance import rows_out, file_resp

CLSID = "{D6E66400-5EED-4943-BEC3-DA55C65D499E}"   # Out_Statement_TLB.pas CLASS_Out_StatementForm
_INTS = re.compile(r"^\s*-?\d+(\s*,\s*-?\d+)*\s*$")
_ACCS = re.compile(r"^[A-Za-z0-9,\s]*$")
_BALS = re.compile(r"^[0-9,\s]*$")


class StIn(BaseModel):
    beg: datetime.date                                # DTP_Begin
    end: datetime.date                                # DTP_End
    turnover: bool = True                             # CBTurnOver (design-time Checked)
    rate_diff: bool = False                           # CBRateDiff
    acc_no: str = Field("", max_length=2000)          # Edit_Acc_No
    bal_cn: str = Field("", max_length=2000)          # Edit_Bal_Cn
    cli_id: str = Field("", max_length=2000)          # EditCli_Id
    currency: int = -1                                # DBLCBavalute.KeyValue (-1 unless CBCurrency checked)
    branch: int = -1                                  # DBLCBfilial.KeyValue
    user: int = -1                                    # dblcbUser.KeyValue
    print_mode: int = Field(0, ge=0, le=1)            # RGroupPrintMode (1 = ForceNewPage per account)
    setting: int = Field(0, ge=0, le=1)               # RGroupSetting (0 Out_RepR, 1 Out_Rep)
    val_view: int = Field(0, ge=0, le=1)              # RGroupValView (0 PairCurrencyId, 1 aValname) - Out_RepR only
    labels: dict[str, str] = Field(default_factory=dict)


def update_str(inp: str) -> str:
    """UpdateStr Out_StatementImpl.pas:382-431, loop for loop (1-based Delphi indices)."""
    result = inp
    if result == "":
        return "' '"
    temp, end_s, beg_index, temp_count, flag = "", "", 1, 0, 1
    temp = "".join(ch for ch in result if ch != " ")
    inp = temp
    result = inp
    temp = ""
    for i in range(1, len(result) + 1):
        if result[i - 1] == ",":
            count = (i - 1) - temp_count
            result = result[beg_index - 1:beg_index - 1 + count]
            temp = temp + result + "'" + "," + "'"
            result = inp
            beg_index = i + 1
            temp_count = i
            flag = 0
        if flag == 1:
            end_s = end_s + result[i - 1]
        else:
            end_s = ""
            flag = 1
    return "'" + temp + end_s + "'"


def turn_and_ratediff(turn: bool, rate: bool) -> str:
    """:894-901. Pascal 'A=true and B=true' parses as (A = (true and B)) = true, so :898 fires when both boxes are
    equal and :900 ('A=false and B=true' -> A = false) whenever CBTurnOver is unchecked."""
    v = ""
    if turn:
        v = "ch_turn"
    if rate:
        v = "ch_rate"
    if turn == (True and rate):
        v = "ch_rate_and_turn"
    if turn == (False and rate):
        v = "ch_nothink"
    return v


def params(b: StIn):
    """:904-928 (+ validation of what the package concatenates)."""
    if b.bal_cn and not _BALS.match(b.bal_cn) or b.acc_no and not _ACCS.match(b.acc_no):
        raise ValueError("lists")
    if b.cli_id != "" and not _INTS.match(b.cli_id):
        raise ValueError("lists")
    bal = "-1" if b.bal_cn == "" else update_str(b.bal_cn)
    acc = "'-1'" if b.acc_no == "" else update_str(b.acc_no)
    cli = "-1" if b.cli_id == "" else b.cli_id
    bd = datetime.datetime.combine(b.beg, datetime.time())             # Integer(Trunc(DTP_Begin.DateTime))
    ed = datetime.datetime.combine(b.end, datetime.time())
    return bd, ed, cli, acc, bal


def export_data_table(c, b: StIn) -> dict:
    trd = turn_and_ratediff(b.turnover, b.rate_diff)
    bd, ed, cli, acc, bal = params(b)
    v_user, v_val, v_filial = int(b.user), int(b.currency), int(b.branch)
    skip_rate = not b.rate_diff                                                       # :932-935
    cstr = c.call_func("OUT_SB_STATEMENT", "notAccId", [bd, ed, cli, acc, bal, v_filial, v_val, v_user, trd])   # :939
    excl = set()
    if b.turnover:                                                                    # :956-971
        excl = {int(r["ACCOUNTID"]) for r in c.query(cstr)}
    binds = {"bDate": bd, "eDate": ed, "ClientID": cli, "AccountNo": acc, "BalanceNo": bal,
             "BranchID": v_filial, "CurrencyID": v_val, "UserID": v_user}
    astr = c.query(Q.SQL_OST_ODS_GETEXTRACTION, binds)[0]["STR"]                     # :992-1004
    bstr = c.query(Q.SQL_OST_ODS_EXTRACTIONTRANS, binds)[0]["STR"]                   # :1018-1030
    main = [r for r in c.query(astr) if int(r["ACCOUNTID"]) not in excl]            # :1037-1048 Filter:=FilterExt
    for m in main:                                                                    # ODSExtractTrans detail (:1049-1060)
        tr = c.query(bstr, {"OwnAccountId": m["ACCOUNTID"]})
        m["TRANS"] = [t for t in tr if not (skip_rate and t.get("OPERATIONCODE") is not None and int(t["OPERATIONCODE"]) == 1)]
        re_ = c.query(Q.SQL_OST_RATESDATASET, {"Val_Kod": m["CURRENCYID"], "Dat": m.get("ENDDATE")}) if m.get("ENDDATE") else []
        rb = c.query(Q.SQL_OST_RATESDATASETBEGIN, {"Val_Kod": m["CURRENCYID"], "Dat": m.get("DATDATE")}) if m.get("DATDATE") else []
        m["RATE_END"] = re_[0]["OSTAT"] if re_ else None
        m["RATE_BEG"] = rb[0]["OSTAT"] if rb else None
    return {"turn_and_ratediff": trd, "not_acc_sql": cstr, "main_sql": astr, "trans_sql": bstr, "excluded": sorted(excl), "accounts": main}


def _f(v) -> float:
    return float(v) if v is not None else 0.0


def report_flags(m: dict) -> dict:
    """Out_RepR/Out_Rep QRSubDetailAfterPrint + group footer sums (SUM(IF(CURRENCYID=0,0,debit)) ...)."""
    cur0 = _f(m["CURRENCYID"]) == 0
    tr = m["TRANS"]
    s = {"ldebit": sum(_f(t["LDEBIT"]) for t in tr), "lkredit": sum(_f(t["LKREDIT"]) for t in tr),
         "debit": 0.0 if cur0 else sum(_f(t["DEBIT"]) for t in tr), "kredit": 0.0 if cur0 else sum(_f(t["KREDIT"]) for t in tr)}
    show = not (s["debit"] == 0 and s["kredit"] == 0) and not (s["ldebit"] == s["debit"] and s["lkredit"] == s["kredit"])
    fx = [not (_f(t["DEBIT"]) == 0 and _f(t["KREDIT"]) == 0) and not (_f(t["PAIRCURRENCYID"]) == 0 and _f(t["CURRENCYID"]) == 0)
          for t in tr]
    return {"sums": s, "show_fx_sums": show, "fx": fx, "count": len(tr), "national": cur0}


def money(v) -> str:
    """FormatFloat('#,##0.00', v)."""
    return f"{Decimal(str(v or 0)):,.2f}"


def rate4(v) -> str:
    return f"{Decimal(str(v or 0)):,.4f}"


def ddmmyyyy(v) -> str:
    return v.strftime("%d.%m.%Y") if v else ""


# Captions used by the files; defaults are the Delphi texts (Azerbaijani), the web sends its locale.
DEFAULT_LABELS = {
    "title": "Hesabdan çıxarış", "acc_no": "Hesab №", "acc_name": "Hesabın adı :", "voen": "VÖEN :", "currency": "Valyuta :",
    "period": "Əməliyyatların tarixi :", "last_op": "Əvvəlki əməliyyat tarixi : ", "begin_rest": "Günün əvvəlinə qalıq :",
    "end_rest": "Günün sonuna qalıq :", "count": "Sənədlərin sayı :", "date": "Tarix", "val": "Val.", "pair": "Qarşı hesab",
    "debit": "Debet", "credit": "Kredit", "purpose": "Təyinatı", "manat": "Manat: ", "valyuta": "Valyuta: ", "rate": "Rəsmi məzənnə: ",
}
# Excel captions are literals in ToExcelClick :556-806
XL = {"acc": " Hesab № :", "name0": "Hesabın adı :", "name1": "Hesabın adı:", "val0": "Valyuta :", "val1": "Valyuta:",
      "opdate": "Əməliy. tarixi:", "last": "Əvv əm. tarixi: ", "beg": "Günün əvv. qal:", "manat": "Manat :", "manat_end": "Manat:",
      "rate": "Rəsmi məzənnə:", "date": "Tarix", "cur": "Valyuta", "pair": "Qarşı hesab", "debit": "Debet", "credit": "Kredit",
      "purpose": "Təyinatı", "end": "Günün son. qal:"}


def to_xlsx(data: dict, b: StIn, lab: dict) -> bytes:
    """ToExcelClick :519-833, cell for cell (Sheet.Cells[row, col])."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side
    L = {**XL, **{k[3:]: v for k, v in lab.items() if k.startswith("xl_")}}
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for col, w in (("A", 15), ("B", 15), ("C", 10), ("D", 20), ("E", 12), ("F", 12), ("G", 100)):
        ws.column_dimensions[col].width = w
    bold = Font(name="Tahoma", bold=True)
    blue = Font(name="Tahoma", bold=True, color="0000FF")
    red = Side(style="thin", color="FF0000")

    def put(r, c, v, font=None, fmt=None):
        cell = ws.cell(row=r, column=c, value=v)
        if font:
            cell.font = font
        if fmt:
            cell.number_format = fmt
        elif c == 4 and isinstance(v, (int, float)):
            cell.number_format = "#######"
        return cell

    def num(r, c, v, f="#,##0.00"):
        put(r, c, round(float(v or 0), 4 if "0000" in f else 2), fmt=f)

    def frame(r1, r2, fill, sep):
        for r in range(r1, r2 + 1):
            for c in range(1, 8):
                cell = ws.cell(row=r, column=c)
                cell.border = Border(left=red, right=red, top=red, bottom=red)
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
        for c in range(1, 8):
            ws.cell(row=sep, column=c).fill = PatternFill("solid", fgColor="00FFFF")   # ColorIndex 8

    def head(r):
        for c, k in ((2, "date"), (3, "cur"), (4, "pair"), (5, "debit"), (6, "credit"), (7, "purpose")):
            put(r, c, L[k], blue)

    accs = data["accounts"]
    d, f, dd = 6, 3, 2
    k = 0
    t_val = int(_f(accs[0]["CURRENCYID"])) if accs else 0               # :550 first record decides the layout
    for i, m in enumerate(accs, 1):
        r0 = i + k
        put(r0, 3, L["acc"], bold); put(r0, 4, m["ACCOUNTNUMBER"])
        put(r0 + 1, 1, L["name0"] if t_val == 0 else L["name1"], bold); put(r0 + 1, 2, m["ACCOUNTNAME"], Font(name="Tahoma"))
        put(r0 + 2, 1, L["val0"] if t_val == 0 else L["val1"], bold); put(r0 + 2, 2, m["CURRENCYNAME"])
        put(r0 + 3, 1, L["opdate"], bold); put(r0 + 3, 2, datetime.datetime.combine(b.beg, datetime.time()), fmt="dd.mm.yyyy")
        put(r0 + 3, 3, datetime.datetime.combine(b.end, datetime.time()), fmt="dd.mm.yyyy")
        put(r0 + 4, 1, L["last"], bold); put(r0 + 4, 2, m.get("LASTOPDATE"), fmt="dd.mm.yyyy")
        put(r0 + 5, 1, L["beg"], bold)
        tr = m["TRANS"]
        mm = len(tr)
        if t_val == 0:                                                   # :551-662
            put(r0 + 5, 2, L["manat"]); num(r0 + 5, 3, abs(_f(m["OSTAT"])))
            head(r0 + 6)
            ld = lk = 0.0
            for j, t in enumerate(tr, 1):
                ld += _f(t["LDEBIT"]); lk += _f(t["LKREDIT"])
                r = i + k + j + d
                put(r, 2, t["DAT"], fmt="dd.mm.yyyy"); put(r, 3, str(t["PAIRCURRENCYID"])); put(r, 4, t["PAIRACCOUNTNAME"])
                num(r, 5, t["LDEBIT"]); num(r, 6, t["LKREDIT"]); put(r, 7, t["NAME"], Font(name="Tahoma"))
            num(i + k + mm + d + 1, 5, ld); num(i + k + mm + d + 1, 6, lk)
            put(i + k + mm + d + 2, 1, L["end"], bold); put(i + k + mm + d + 2, 2, L["manat_end"])
            num(i + k + mm + d + 2, 3, abs(_f(m["ENDOSTAT"])))
            frame(r0, i + k + mm + d + 2, None, i + k + mm + d + 3)
            k = k + d + f + mm
        else:                                                            # :665-817
            put(r0 + 5, 2, L["val1"], Font(name="Tahoma")); num(r0 + 5, 3, abs(_f(m["OSTAT"])))
            put(r0 + 6, 2, L["manat"], Font(name="Tahoma")); num(r0 + 6, 3, abs(_f(m["POKOSTAT"])))
            put(r0 + 7, 2, L["rate"], Font(name="Tahoma")); num(r0 + 7, 3, abs(_f(m["RATE_BEG"])), "#,##0.0000")
            head(r0 + 8)
            www = 0
            ld = lk = de = kr = 0.0
            for j, t in enumerate(tr, 1):
                de += _f(t["DEBIT"]); kr += _f(t["KREDIT"]); ld += _f(t["LDEBIT"]); lk += _f(t["LKREDIT"])
                r = i + k + j + d + dd + www
                put(r, 2, t["DAT"], fmt="dd.mm.yyyy"); put(r, 3, str(t["PAIRCURRENCYID"])); put(r, 4, t["PAIRACCOUNTNAME"])
                num(r, 5, t["LDEBIT"]); num(r, 6, t["LKREDIT"]); put(r, 7, t["NAME"], Font(name="Tahoma"))
                num(r + 1, 5, t["DEBIT"]); num(r + 1, 6, t["KREDIT"])
                nxt = tr[min(j, mm - 1)]                                 # :771 Next, then the (new / last) record is tested
                if int(_f(nxt["PAIRCURRENCYID"])) != 0:
                    www += 1
                    if j == mm and www != 0:
                        www += 1
            base = i + k + mm + d + dd + www
            num(base, 5, ld); num(base, 6, lk)
            num(base + 1, 5, de); num(base + 1, 6, kr)
            put(base + 2, 1, L["end"], bold); put(base + 2, 2, L["val1"]); num(base + 2, 3, abs(_f(m["ENDOSTAT"])))
            put(base + 3, 2, L["manat"]); num(base + 3, 3, abs(_f(m["ENDPOKOSTAT"])))
            put(base + 4, 2, L["rate"]); num(base + 4, 3, abs(_f(m["RATE_END"])), "#,##0.0000")
            frame(r0, base + 4, "FFFF00", base + 5)                     # ColorIndex 6 / 8
            k = k + d + 2 * dd + www + f + mm
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _period(b: StIn) -> str:
    """QRD_Begin / QRD_End; QRBandDetailBeforePrint blanks QRD_End when both captions are equal."""
    s, e = b.beg.strftime("%d.%m.%Y"), b.end.strftime("%d.%m.%Y")
    return s if s == e else f"{s}  {e}"


def to_txt(data: dict, b: StIn, lab: dict, bank: str) -> bytes:
    """PostClick :834-849: FormRep (Out_Rep, purpose below) through TQRAsciiExportFilter - plain text, one block per account."""
    L = {**DEFAULT_LABELS, **lab}
    out = []
    for m in data["accounts"]:
        fl = report_flags(m)
        nat = fl["national"]
        out += [bank, f"{L['acc_no']} {m['ACCOUNTNUMBER']}", f"{L['acc_name']} {m['ACCOUNTNAME']}",
                f"{L['voen']} {m.get('CUSTOMERVOIN') or ''}    {L['currency']} {m['CURRENCYNAME']}",
                f"{L['period']} {_period(b)}", f"{L['last_op']} {ddmmyyyy(m.get('LASTOPDATE'))}  {ddmmyyyy(m.get('OPENDATE'))}",
                f"{L['begin_rest']} {L['manat'] if nat else L['valyuta']}{money(abs(_f(m['OSTAT'])))} {m.get('AP_BEG') or ''}"
                + ("" if nat else f"  {L['manat']}{money(abs(_f(m['POKOSTAT'])))}  {L['rate']}{rate4(abs(_f(m['RATE_BEG'])))}"),
                "-" * 100, f"{L['date']:<12}{L['val']:<6}{L['pair']:<30}{L['debit']:>18}{L['credit']:>18}", "-" * 100]
        for t, fx in zip(m["TRANS"], fl["fx"]):
            out.append(f"{ddmmyyyy(t['DAT']):<12}{str(t['PAIRCURRENCYID']):<6}{str(t['PAIRACCOUNTNAME'] or ''):<30}"
                       f"{money(t['LDEBIT']):>18}{money(t['LKREDIT']):>18}")
            if fx:
                out.append(" " * 48 + f"{money(t['DEBIT']):>18}{money(t['KREDIT']):>18}")
            out.append(f"    {L['purpose']}: {t['NAME'] or ''}")
        s = fl["sums"]
        out += ["-" * 100, " " * 48 + f"{money(s['ldebit']):>18}{money(s['lkredit']):>18}"]
        if fl["show_fx_sums"]:
            out.append(" " * 48 + f"{money(s['debit']):>18}{money(s['kredit']):>18}")
        out += [f"{L['end_rest']} {L['manat'] if nat else L['valyuta']}{money(abs(_f(m['ENDOSTAT'])))} {m.get('AP_END') or ''}"
                + ("" if nat else f"  {L['manat']}{money(abs(_f(m['ENDPOKOSTAT'])))}  {L['rate']}{rate4(abs(_f(m['RATE_END'])))}"),
                f"{L['count']} {fl['count']}", "", ""]
    return "\r\n".join(out).encode("utf-8")


def to_pdf(data: dict, b: StIn, lab: dict, bank: str) -> bytes:
    """PreviewClick :852-861 -> ShowReportR (:477, RGroupSetting 0) or ShowReport (:458, 1); RGroupPrintMode 1 = page per account."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak
    from xml.sax.saxutils import escape
    L = {**DEFAULT_LABELS, **lab}
    f = R._font()
    fb = f + "-Bold" if f != "Helvetica" else "Helvetica-Bold"
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=A4, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
                            title=L["title"])
    st = ParagraphStyle("h", fontName=f, fontSize=8.5, leading=11)
    stb = ParagraphStyle("b", fontName=fb, fontSize=10, leading=13, textColor=colors.HexColor("#0B2545"))
    right = b.setting == 0
    story = []
    for n, m in enumerate(data["accounts"]):
        if n and b.print_mode == 1:
            story.append(PageBreak())
        fl = report_flags(m)
        nat = fl["national"]
        P = lambda t, s=st: Paragraph(escape(str(t)), s)   # noqa: E731
        story += [P(bank, stb), P(f"{L['acc_no']} {m['ACCOUNTNUMBER']}", stb), P(f"{L['acc_name']} {m['ACCOUNTNAME']}"),
                  P(f"{L['voen']} {m.get('CUSTOMERVOIN') or ''}     {L['currency']} {m['CURRENCYNAME']}"),
                  P(f"{L['period']} {_period(b)}"),
                  P(f"{L['last_op']} {ddmmyyyy(m.get('LASTOPDATE'))}   {ddmmyyyy(m.get('OPENDATE'))}"),
                  P(f"{L['begin_rest']} {L['manat'] if nat else L['valyuta']}({money(abs(_f(m['OSTAT'])))}) {m.get('AP_BEG') or ''}"
                    + ("" if nat else f"   {L['manat']}{money(abs(_f(m['POKOSTAT'])))}   {L['rate']}{rate4(abs(_f(m['RATE_BEG'])))}")),
                  Spacer(1, 2 * mm)]
        cur_field = "AVALNAME" if (right and b.val_view == 1) else "PAIRCURRENCYID"     # ShowReportR :488-491
        hdr = [L["date"], L["val"], L["pair"], L["debit"], L["credit"]] + ([L["purpose"]] if right else [])
        rows, styles = [hdr], [("FONTNAME", (0, 0), (-1, -1), f), ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B2545")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("ALIGN", (3, 1), (4, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#9DB4C0")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for t, fx in zip(m["TRANS"], fl["fx"]):
            base = [ddmmyyyy(t["DAT"]), str(t.get(cur_field) or ""), str(t["PAIRACCOUNTNAME"] or ""), money(t["LDEBIT"]), money(t["LKREDIT"])]
            rows.append(base + ([Paragraph(escape(t["NAME"] or ""), st)] if right else []))
            if fx:                                                           # ChBand / DBTextDebit+DBTextKredit
                rows.append(["", "", "", money(t["DEBIT"]), money(t["KREDIT"])] + ([""] if right else []))
            if not right:
                rows.append([Paragraph(escape(f"{L['purpose']}: {t['NAME'] or ''}"), st), "", "", "", ""])
                styles.append(("SPAN", (0, len(rows) - 1), (-1, len(rows) - 1)))
        s = fl["sums"]
        rows.append(["", "", "", money(s["ldebit"]), money(s["lkredit"])] + ([""] if right else []))
        styles.append(("FONTNAME", (0, len(rows) - 1), (-1, len(rows) - 1), fb))
        if fl["show_fx_sums"]:
            rows.append(["", "", "", money(s["debit"]), money(s["kredit"])] + ([""] if right else []))
            styles.append(("FONTNAME", (0, len(rows) - 1), (-1, len(rows) - 1), fb))
        widths = [22 * mm, 14 * mm, 48 * mm, 26 * mm, 26 * mm] + ([54 * mm] if right else [])
        if not right:
            widths[2] = 102 * mm
        t = Table(rows, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle(styles))
        story += [t, Spacer(1, 2 * mm),
                  P(f"{L['end_rest']} {L['manat'] if nat else L['valyuta']}({money(abs(_f(m['ENDOSTAT'])))}) {m.get('AP_END') or ''}"
                    + ("" if nat else f"   {L['manat']}{money(abs(_f(m['ENDPOKOSTAT'])))}   {L['rate']}{rate4(abs(_f(m['RATE_END'])))}")),
                  P(f"{L['count']} {fl['count']}"), Spacer(1, 8 * mm)]
    if not story:
        story = [Paragraph(escape(L["title"]), stb)]
    doc.build(story)
    return bio.getvalue()


def make_router(current, require_module, err):
    r = APIRouter(prefix="/api/out-statement")
    g = require_module(CLSID)

    @r.get("/init")
    def _init(ws=Depends(g)):
        c = ws.conn                                                     # ActiveFormCreate :1082-1103 + active lookups
        day = c.query(Q.SQL_OST_ODSSYSDAT)[0]["TEK_DAT"]
        day = day.date() if isinstance(day, datetime.datetime) else day
        return {"date": day.isoformat(),
                "currencies": [{"id": x["VALUTE"], "name": x["VALU_NAME"]} for x in c.query(Q.SQL_OST_ODSCURRENCY)],
                "branches": [{"id": x["ID"], "name": x["NAME"]} for x in c.query(Q.SQL_OST_ODSFILIAL)],
                "users": [{"id": x["ID"], "name": x["NAME"]} for x in c.query(Q.SQL_OST_ODSUSERS)]}

    def data(c, b):
        try:
            return export_data_table(c, b)
        except ValueError:
            err(400, "out_statement.err_lists")

    def bank(c):
        rows = c.query(Q.SQL_OST_ODSLOGO)                               # QRDBTextBank AZ_NAME
        return (rows[0].get("AZ_NAME") or "") if rows else ""

    @r.post("/run")
    def _run(b: StIn, ws=Depends(g)):
        d = data(ws.conn, b)
        accs = []
        for m in d["accounts"]:
            fl = report_flags(m)
            x = rows_out([{k: v for k, v in m.items() if k != "TRANS"}])[0]
            x["trans"] = rows_out(m["TRANS"])
            for t, fx in zip(x["trans"], fl["fx"]):
                t["fx"] = fx
            x["sums"], x["show_fx_sums"], x["count"] = fl["sums"], fl["show_fx_sums"], fl["count"]
            accs.append(x)
        return {"turn_and_ratediff": d["turn_and_ratediff"], "excluded": d["excluded"], "accounts": accs, "bank": bank(ws.conn)}

    @r.post("/statement.{fmt}")
    def _file(fmt: str, b: StIn, ws=Depends(g)):
        if fmt not in ("pdf", "xlsx", "txt"):
            err(404, "error.generic")
        d = data(ws.conn, b)
        lab = {k: str(v)[:200] for k, v in b.labels.items()}
        if fmt == "xlsx":
            return file_resp(to_xlsx(d, b, lab), "xlsx", "out-statement")
        if fmt == "txt":
            return Response(to_txt(d, b, lab, bank(ws.conn)), media_type="text/plain; charset=utf-8",
                            headers={"Content-Disposition": 'attachment; filename="out-statement.txt"'})
        return file_resp(to_pdf(d, b, lab, bank(ws.conn)), "pdf", "out-statement")

    return r
