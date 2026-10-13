"""Cards step 1 E2E + screenshots: Packages/Cards_New/cards_azericard main menu + 'Müştərilər' list."""
import os, sys
from playwright.sync_api import sync_playwright
BASE = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:3000"
OUT = os.path.join(os.path.dirname(__file__), "..", "docs", "screenshots")
CHROME = os.getenv("CHROME_PATH", "/usr/bin/google-chrome")
shot = lambda pg, n: (pg.screenshot(path=os.path.join(OUT, n), full_page=False), print("saved", n))
with sync_playwright() as p:
    b = p.chromium.launch(executable_path=CHROME if os.path.exists(CHROME) else None, args=["--no-sandbox"])
    pg = b.new_page(viewport={"width": 1440, "height": 1000})
    pg.on("pageerror", lambda e: print("pageerror:", e))
    pg.goto(BASE); pg.wait_for_selector("#loginForm")
    pg.fill("#u", "ADMIN"); pg.fill("#p", "Admin#2026"); pg.click("#loginBtn"); pg.wait_for_selector("#tree .node")
    pg.fill("#treeSearch", "Cards"); pg.wait_for_timeout(300)
    shot(pg, "160-cards-tree-entry.png")
    pg.locator("#tree .node", has_text="Cards").last.click(); pg.wait_for_selector("#g_c tr[data-i='4']")
    assert pg.inner_text("#c_cnt") == "5"
    shot(pg, "161-cards-clients.png")
    pg.click("#m1 summary"); pg.wait_for_timeout(200)
    shot(pg, "162-cards-menu.png")
    pg.click("#m1 summary")
    pg.click("#c_flt"); pg.check("#c_p5 [data-c='2']"); pg.select_option("#c_l2", "AZN"); pg.check("#c_p5 [data-c='4']")
    pg.select_option("#c_l4", "ACTIVE"); pg.click("#c_apply"); pg.wait_for_timeout(500)
    assert pg.inner_text("#c_cnt") == "3", pg.inner_text("#c_cnt")
    shot(pg, "163-cards-clients-filter.png")
    pg.click("#c_srch"); pg.fill("#c_p4 [data-e='FNAME']", "R"); pg.wait_for_timeout(500)
    assert pg.inner_text("#c_cnt") == "1"
    shot(pg, "164-cards-clients-search.png")
    with pg.expect_download() as dl:
        pg.click("#c_xls")
    path = dl.value.path()
    from openpyxl import load_workbook
    import io
    sh = load_workbook(io.BytesIO(open(path, "rb").read())).active
    assert sh.max_column == 28 and sh.cell(1, 2).value == "RBSNUMBER", sh.max_column
    print("excel rows", sh.max_row - 1)
    pg.wait_for_timeout(300)
    shot(pg, "165-cards-clients-excel.png")
    b.close()
